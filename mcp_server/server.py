"""
Context Integrity SoK — MCP server.

Exposes the distilled output of a ~1000-paper systematic review (LLM agent
context poisoning: adversarial attacks vs. incidental degradation) as
queryable tools: papers, named pollution mechanisms (RQ3), confirmed
defenses (RQ4), the defense x mechanism coverage matrix (RQ5), and
plain-language headline findings per research question (RQ1-RQ6).

All data is read from files already committed in this repo (data/exports/,
data/registries/, and the rq*.md writeups) — nothing is fetched over the
network and no API key is required to run this server.
"""
import json
import csv
import glob
import re
from pathlib import Path

from mcp.server.mcpserver import MCPServer
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

# ================================================================ DATA LOADING

def _load_papers():
    wb = load_workbook(ROOT / "data/exports/paper_dashboard_source.xlsx", read_only=True)
    ws = wb["Papers"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    return [dict(zip(header, r)) for r in rows[1:]]


def _load_registries():
    mechs = json.load(open(ROOT / "data/registries/rq3_pollution_registry.json"))
    defs = json.load(open(ROOT / "data/registries/rq4_defense_registry.json"))
    cov = json.load(open(ROOT / "data/registries/rq5_coverage_matrix.json"))
    justif = {}
    for fp in sorted(glob.glob(str(ROOT / "data/registries/raw/rq5_batch*.csv"))):
        with open(fp, newline="") as f:
            for row in csv.DictReader(f):
                key = (str(row.get("defense_row", "")).strip(), row.get("matched_mechanism_name", "").strip())
                justif[key] = row.get("match_justification", "")
    return mechs, defs, cov, justif


PAPERS = _load_papers()
PAPERS_BY_ID = {p["paper_id"]: p for p in PAPERS}
MECHS, DEFS, COV, JUSTIF = _load_registries()
MECH_TO_DEFENSES = COV["mech_to_defenses"]       # mechanism_name -> [defense rows]
DEFENSE_TO_MECHS = COV["defense_to_mechs"]       # defense row (str) -> [mechanism names]
UNCOVERED_NAMES = set(COV["uncovered_mechanisms"])
MECHS_BY_NAME = {(m.get("technique_name") or m.get("name")): m for m in MECHS}
DEFS_BY_ROW = {str(d["row"]): d for d in DEFS}
DEFS_BY_NAME = {d.get("defense_name"): d for d in DEFS}

RQ_FILES = {
    "RQ1": ROOT / "rq1_taxonomy_analysis.md",
    "RQ2": ROOT / "cross_citation_analysis.md",
    "RQ3": ROOT / "rq3_pollution_census.md",
    "RQ4": ROOT / "rq4_defense_census.md",
    "RQ5": ROOT / "rq5_coverage_matrix.md",
    "RQ6": ROOT / "rq6_case_studies.md",
}


def _match(text, needle):
    return needle.lower() in (text or "").lower()


def _paper_summary(p):
    return {
        "paper_id": p.get("paper_id"),
        "title": p.get("title"),
        "year": p.get("year"),
        "venue": p.get("venue"),
        "track": p.get("track"),
        "channel": p.get("channel"),
        "consequence": p.get("consequence"),
        "defense_intervention_point": p.get("defense_intervention_point"),
        "citation_count": p.get("citation_count"),
    }


# ================================================================ SERVER

mcp = MCPServer("context-integrity-sok")


@mcp.tool()
def get_stats() -> dict:
    """Headline counts for the whole Context Integrity SoK corpus: how many
    papers, named pollution mechanisms, confirmed defenses, and confirmed
    (defense, mechanism) tested pairs. Good first call to orient yourself."""
    return {
        "papers_in_corpus": len(PAPERS),
        "named_mechanisms_rq3": len(MECHS),
        "confirmed_defenses_rq4": len(DEFS),
        "confirmed_defense_mechanism_pairs_rq5": COV["n_matched_pairs"],
        "defenses_with_confirmed_mechanism_match": COV["n_defenses_matched"],
        "defenses_with_confirmed_mechanism_match_pct": round(100 * COV["n_defenses_matched"] / COV["n_defenses_total"], 1),
        "mechanisms_with_at_least_one_defense": COV["n_mechs_covered"],
        "mechanisms_with_zero_defenses": len(UNCOVERED_NAMES),
        "mechanisms_with_zero_defenses_pct": round(100 * len(UNCOVERED_NAMES) / COV["n_mechs_total"], 1),
    }


@mcp.tool()
def list_papers(query: str = "", track: str = "", channel: str = "", consequence: str = "",
                defense_intervention_point: str = "", limit: int = 25) -> list[dict]:
    """Search the full 1008-paper corpus. `query` substring-matches title,
    key_result, and technical_summary (case-insensitive). `track` is
    'Security', 'ML/AI', or 'Both'. Returns compact summaries — use
    get_paper(paper_id) for the full extracted record (key_result,
    baselines_compared, technical_summary, etc)."""
    out = []
    for p in PAPERS:
        if track and str(p.get("track", "")).strip().lower() != track.strip().lower():
            continue
        if channel and str(p.get("channel", "")).strip().lower() != channel.strip().lower():
            continue
        if consequence and str(p.get("consequence", "")).strip().lower() != consequence.strip().lower():
            continue
        if defense_intervention_point and str(p.get("defense_intervention_point", "")).strip().lower() != defense_intervention_point.strip().lower():
            continue
        if query and not any(_match(p.get(f), query) for f in ("title", "key_result", "technical_summary")):
            continue
        out.append(_paper_summary(p))
        if len(out) >= limit:
            break
    return out


@mcp.tool()
def get_paper(paper_id: str) -> dict:
    """Full extracted record for one paper by paper_id (from list_papers or
    list_mechanisms/list_defenses' source_paper_id field) — includes
    abstract, key_result, baselines_compared, technical_summary,
    stated_limitations, models_evaluated, datasets_benchmarks, etc."""
    p = PAPERS_BY_ID.get(paper_id)
    if not p:
        return {"error": f"no paper with paper_id={paper_id!r}"}
    return dict(p)


@mcp.tool()
def list_mechanisms(query: str = "", track: str = "", channel: str = "", consequence: str = "",
                     covered_only: bool | None = None, limit: int = 50) -> list[dict]:
    """List named pollution/attack mechanisms from the RQ3 registry (183
    total: both deliberate attack techniques and incidental degradation
    mechanisms). Set covered_only=True for mechanisms with >=1 confirmed
    defense tested against them, or False for the 116 with zero coverage."""
    out = []
    for m in MECHS:
        nm = m.get("technique_name") or m.get("name")
        if track and str(m.get("track", "")).strip().lower() != track.strip().lower():
            continue
        if channel and str(m.get("channel", "")).strip().lower() != channel.strip().lower():
            continue
        if consequence and str(m.get("consequence", "")).strip().lower() != consequence.strip().lower():
            continue
        is_covered = nm not in UNCOVERED_NAMES
        if covered_only is not None and is_covered != covered_only:
            continue
        if query and not any(_match(m.get(f), query) for f in ("technique_name", "name", "notes")):
            continue
        defenders = MECH_TO_DEFENSES.get(nm, [])
        out.append({
            "mechanism_name": nm,
            "track": m.get("track"),
            "channel": m.get("channel"),
            "consequence": m.get("consequence"),
            "n_defenses_tested_against": len(defenders),
            "source_paper_id": m.get("paper_id"),
            "source_paper_title": m.get("title"),
            "notes": m.get("notes"),
        })
        if len(out) >= limit:
            break
    return out


@mcp.tool()
def list_defenses(query: str = "", track: str = "", channel: str = "", consequence: str = "",
                   defense_intervention_point: str = "", validated_against: str = "",
                   has_mechanism_match: bool | None = None, limit: int = 50) -> list[dict]:
    """List confirmed defenses from the RQ4 registry (479 total).
    `defense_intervention_point` is 'ingestion', 'reasoning', or 'execution'.
    `validated_against` is 'adversarial' or 'incidental' (the threat model
    the defense's own paper tested it against). has_mechanism_match=False
    lists the 309 defenses NOT confirmed tested against any RQ3-named
    mechanism (see get_stats / RQ5 for why that isn't the same as "untested")."""
    out = []
    for d in DEFS:
        if track and str(d.get("track", "")).strip().lower() != track.strip().lower():
            continue
        if channel and str(d.get("channel", "")).strip().lower() != channel.strip().lower():
            continue
        if consequence and str(d.get("consequence", "")).strip().lower() != consequence.strip().lower():
            continue
        if defense_intervention_point and str(d.get("defense_intervention_point", "")).strip().lower() != defense_intervention_point.strip().lower():
            continue
        if validated_against and str(d.get("validated_against", "")).strip().lower() != validated_against.strip().lower():
            continue
        mechs_tested = DEFENSE_TO_MECHS.get(str(d["row"]), [])
        if has_mechanism_match is not None and (len(mechs_tested) > 0) != has_mechanism_match:
            continue
        if query and not any(_match(d.get(f), query) for f in ("defense_name", "notes")):
            continue
        out.append({
            "defense_name": d.get("defense_name"),
            "track": d.get("track"),
            "channel": d.get("channel"),
            "consequence": d.get("consequence"),
            "defense_intervention_point": d.get("defense_intervention_point"),
            "validated_against": d.get("validated_against"),
            "n_mechanisms_tested_against": len(mechs_tested),
            "mechanisms_tested_against": mechs_tested,
            "source_paper_id": d.get("paper_id"),
            "source_paper_title": d.get("title"),
            "notes": d.get("notes"),
        })
        if len(out) >= limit:
            break
    return out


@mcp.tool()
def get_coverage_for_mechanism(mechanism_name: str) -> dict:
    """All confirmed defenses tested against one named mechanism (exact
    name, as returned by list_mechanisms), with the extracting agent's
    justification text for each match."""
    if mechanism_name not in MECHS_BY_NAME:
        return {"error": f"no mechanism named {mechanism_name!r} — call list_mechanisms(query=...) to find the exact name"}
    rows = MECH_TO_DEFENSES.get(mechanism_name, [])
    defenders = []
    for row in rows:
        d = DEFS_BY_ROW.get(str(row), {})
        defenders.append({
            "defense_name": d.get("defense_name"),
            "defense_intervention_point": d.get("defense_intervention_point"),
            "validated_against": d.get("validated_against"),
            "justification": JUSTIF.get((str(row), mechanism_name), ""),
        })
    return {"mechanism_name": mechanism_name, "n_defenses": len(defenders), "defenses": defenders}


@mcp.tool()
def get_coverage_for_defense(defense_name: str) -> dict:
    """All mechanisms one named defense (exact name, as returned by
    list_defenses) was confirmed tested against, with justification text."""
    d = DEFS_BY_NAME.get(defense_name)
    if not d:
        return {"error": f"no defense named {defense_name!r} — call list_defenses(query=...) to find the exact name"}
    mechs = DEFENSE_TO_MECHS.get(str(d["row"]), [])
    tested = [{"mechanism_name": m, "justification": JUSTIF.get((str(d["row"]), m), "")} for m in mechs]
    return {"defense_name": defense_name, "n_mechanisms": len(tested), "mechanisms": tested}


@mcp.tool()
def list_uncovered_mechanisms(track: str = "", channel: str = "", limit: int = 200) -> list[dict]:
    """The 116 named mechanisms with ZERO confirmed defenses tested against
    them — a direct 'what's left to defend' list, useful for scoping new
    defense work or a related-work gap statement."""
    out = []
    for nm in sorted(UNCOVERED_NAMES):
        m = MECHS_BY_NAME.get(nm, {})
        if track and str(m.get("track", "")).strip().lower() != track.strip().lower():
            continue
        if channel and str(m.get("channel", "")).strip().lower() != channel.strip().lower():
            continue
        out.append({
            "mechanism_name": nm,
            "track": m.get("track"),
            "channel": m.get("channel"),
            "consequence": m.get("consequence"),
            "source_paper_title": m.get("title"),
        })
        if len(out) >= limit:
            break
    return out


@mcp.tool()
def get_rq_summary(rq: str) -> str:
    """Plain-language headline finding for one research question. `rq` is
    one of RQ1 (taxonomy), RQ2 (citation network), RQ3 (pollution census),
    RQ4 (defense census), RQ5 (coverage matrix), RQ6 (defense
    generalization case studies)."""
    key = rq.strip().upper()
    fp = RQ_FILES.get(key)
    if not fp or not fp.exists():
        return f"Unknown or missing RQ: {rq!r}. Valid values: {', '.join(RQ_FILES)}"
    text = fp.read_text()
    m = re.search(r"^## Headline result\s*\n(.*?)(?=\n## |\Z)", text, re.S | re.M)
    return m.group(1).strip() if m else text[:2000]


if __name__ == "__main__":
    mcp.run()
