import json

import pytest

from src import db, excel_source


def _seed_paper(conn, **overrides):
    base = {
        "paper_id": "p1", "title": "Test Paper", "abstract": "An abstract.",
        "authors": ["A. One", "B. Two"], "year": 2024, "venue": "arXiv.org",
        "doi": None, "arxiv_id": "2401.00001", "url": None, "citation_count": 12,
        "pdf_local_path": None, "track_auto": "A", "track_human": None,
        "screen_auto": "auto_include", "screen_human": None, "discovered_via": "seed",
        "hop": 0, "source_paper_id": None, "discovered_query": None,
        "seed_category": "seed", "date_added": "2026-01-01T00:00:00",
    }
    base.update(overrides)
    db.upsert_paper(conn, base)
    conn.commit()


@pytest.fixture
def conn():
    c = db.connect(":memory:")
    db.init_db(c)
    return c


def test_build_and_reload_roundtrip(tmp_path, conn):
    _seed_paper(conn)
    path = tmp_path / "papers.xlsx"
    excel_source.export_to_path(conn, path)

    df = excel_source.load_dataframe(path)
    assert len(df) == 1
    row = df.iloc[0]
    assert row["paper_id"] == "p1"
    assert row["effective_track"] == "A"  # "Security" display label parsed back
    assert row["effective_screen"] == "auto_include"  # "Include" parsed back
    assert row["venue_short"] == "arXiv"  # normalized from "arXiv.org"
    assert row["authors"] == ["A. One", "B. Two"]


def test_hand_edited_track_and_screening_are_respected(tmp_path, conn):
    _seed_paper(conn)
    path = tmp_path / "papers.xlsx"
    excel_source.export_to_path(conn, path)

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["Papers"]
    header = [c.value for c in ws[1]]
    track_col = header.index("track") + 1
    screen_col = header.index("screening") + 1
    ws.cell(row=2, column=track_col, value="ML/AI")
    ws.cell(row=2, column=screen_col, value="Exclude")
    wb.save(path)

    df = excel_source.load_dataframe(path)
    assert df.iloc[0]["effective_track"] == "B"
    assert df.iloc[0]["effective_screen"] == "auto_exclude"


def test_manually_added_row_gets_a_stable_generated_id(tmp_path, conn):
    _seed_paper(conn)
    path = tmp_path / "papers.xlsx"
    excel_source.export_to_path(conn, path)

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["Papers"]
    header = [c.value for c in ws[1]]
    title_col = header.index("title") + 1
    ws.cell(row=3, column=title_col, value="A Brand New Paper Found By Hand")
    wb.save(path)

    df1 = excel_source.load_dataframe(path)
    df2 = excel_source.load_dataframe(path)
    new_row1 = df1[df1["title"] == "A Brand New Paper Found By Hand"].iloc[0]
    new_row2 = df2[df2["title"] == "A Brand New Paper Found By Hand"].iloc[0]
    assert new_row1["paper_id"] == new_row2["paper_id"]  # stable across reloads
    assert new_row1["paper_id"].startswith("manual:")


def test_export_refuses_to_overwrite_without_flag(tmp_path, conn):
    _seed_paper(conn)
    path = tmp_path / "papers.xlsx"
    excel_source.export_to_path(conn, path)
    with pytest.raises(FileExistsError):
        excel_source.export_to_path(conn, path)
    excel_source.export_to_path(conn, path, overwrite=True)  # should not raise


def test_export_respects_id_filter(tmp_path, conn):
    _seed_paper(conn, paper_id="p1")
    _seed_paper(conn, paper_id="p2", title="Second Paper")
    path = tmp_path / "papers.xlsx"
    excel_source.export_to_path(conn, path, ids={"p1"})

    df = excel_source.load_dataframe(path)
    assert len(df) == 1
    assert df.iloc[0]["paper_id"] == "p1"
