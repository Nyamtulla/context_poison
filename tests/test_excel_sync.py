from unittest.mock import patch

import openpyxl
import pytest

from src import db, excel_source, excel_sync
from src.config import Config

RAW_CONFIG = {
    "search": {"start_year": 2023, "end_year": None},
    "hop_depth": 1,
    "tracks": {"A": {"keyword_clusters": []}, "B": {"keyword_clusters": []}},
    "venues": {"soft_list": []},
    "screening": {
        "include_signal_terms": ["agent", "RAG", "memory"],
        "jailbreak_only_terms": ["jailbreak"],
        "training_time_only_terms": ["training data poisoning"],
        "runtime_signal_terms": ["inference time"],
        "llm_presence_terms": ["LLM", "GPT"],
    },
    "track_signatures": {"A": ["Greshake"], "B": ["Lost in the Middle"]},
    "apis": {
        "semantic_scholar": {
            "base_url": "https://api.semanticscholar.org/graph/v1", "api_key_env": "S2_API_KEY",
            "rate_limit_rps_unauthenticated": 100, "rate_limit_rps_authenticated": 100, "max_retries": 1,
        },
        "arxiv": {
            "base_url": "http://export.arxiv.org/api/query", "pdf_base_url": "https://arxiv.org/pdf",
            "rate_limit_rps": 100, "max_retries": 1,
        },
        "openalex": {
            "base_url": "https://api.openalex.org", "contact_email_env": "OPENALEX_CONTACT_EMAIL",
            "rate_limit_rps": 100, "max_retries": 1,
        },
    },
    "paths": {"raw_cache_dir": "data/cache/raw_responses"},
}


@pytest.fixture
def config(tmp_path):
    raw = dict(RAW_CONFIG)
    raw["paths"] = {"raw_cache_dir": str(tmp_path / "cache")}
    return Config(raw=raw, repo_root=tmp_path)


@pytest.fixture
def conn():
    c = db.connect(":memory:")
    db.init_db(c)
    return c


def _make_sheet(tmp_path, rows):
    conn0 = db.connect(":memory:")
    db.init_db(conn0)
    path = tmp_path / "papers.xlsx"
    excel_source.export_to_path(conn0, path)  # empty, just gets header row
    wb = openpyxl.load_workbook(path)
    ws = wb[excel_source.SHEET_NAME]
    header = [c.value for c in ws[1]]
    for row in rows:
        ws.append([row.get(col, "") for col in header])
    wb.save(path)
    return path


def test_new_paper_with_arxiv_id_gets_resolved_and_synced(tmp_path, config, conn):
    path = _make_sheet(tmp_path, [{
        "title": "", "arxiv_id": "2302.12173", "track": "Security", "screening": "Include",
    }])

    fake_paper = {
        "paper_id": "s2-abc123", "title": "Not What You've Signed Up For", "abstract": "Indirect prompt injection.",
        "authors": ["Kai Greshake"], "year": 2023, "venue": "AISec@CCS", "doi": None,
        "arxiv_id": "2302.12173", "url": None, "citation_count": 1639,
    }
    with patch("src.clients.semantic_scholar.SemanticScholarClient.get_paper", return_value=fake_paper), \
         patch("src.clients.semantic_scholar.SemanticScholarClient.get_references", return_value={"papers": [], "raw_pages": [], "hit_count": 0}), \
         patch("src.clients.semantic_scholar.SemanticScholarClient.get_citations", return_value={"papers": [], "raw_pages": [], "hit_count": 0}):
        counts = excel_sync.sync_excel(conn, config, path)

    assert counts["newly_synced"] == 1
    assert counts["already_synced"] == 0

    row = db.get_paper(conn, "s2-abc123")
    assert row is not None
    assert row["title"] == "Not What You've Signed Up For"
    assert row["track_human"] == "A"  # "Security" resolved back to the internal code
    assert row["screen_human"] == "auto_include"  # "Include" resolved back
    assert row["discovered_via"] == "manual"

    wb = openpyxl.load_workbook(path)
    ws = wb[excel_source.SHEET_NAME]
    header = [c.value for c in ws[1]]
    pid_col = header.index("paper_id") + 1
    abstract_col = header.index("abstract") + 1
    assert ws.cell(row=2, column=pid_col).value == "s2-abc123"
    assert ws.cell(row=2, column=abstract_col).value == "Indirect prompt injection."


def test_rerun_skips_already_synced_rows(tmp_path, config, conn):
    path = _make_sheet(tmp_path, [{"title": "", "arxiv_id": "2302.12173", "track": "Security", "screening": "Include"}])
    fake_paper = {
        "paper_id": "s2-abc123", "title": "T", "abstract": "A", "authors": [], "year": 2023,
        "venue": "V", "doi": None, "arxiv_id": "2302.12173", "url": None, "citation_count": 5,
    }
    with patch("src.clients.semantic_scholar.SemanticScholarClient.get_paper", return_value=fake_paper), \
         patch("src.clients.semantic_scholar.SemanticScholarClient.get_references", return_value={"papers": [], "raw_pages": [], "hit_count": 0}), \
         patch("src.clients.semantic_scholar.SemanticScholarClient.get_citations", return_value={"papers": [], "raw_pages": [], "hit_count": 0}):
        excel_sync.sync_excel(conn, config, path)
        counts2 = excel_sync.sync_excel(conn, config, path)

    assert counts2["already_synced"] == 1
    assert counts2["newly_synced"] == 0
    assert db.paper_count(conn) == 1


def test_unresolvable_paper_falls_back_to_manual_id_without_crashing(tmp_path, config, conn):
    path = _make_sheet(tmp_path, [{"title": "Some Totally Obscure Unfindable Paper", "track": "", "screening": ""}])
    with patch("src.clients.semantic_scholar.SemanticScholarClient.search", return_value={"papers": [], "raw_pages": [], "hit_count": 0}):
        counts = excel_sync.sync_excel(conn, config, path)

    assert counts["newly_synced"] == 1
    rows = db.all_papers(conn)
    assert len(rows) == 1
    assert rows[0]["paper_id"].startswith("manual:")
    assert rows[0]["track_human"] is None  # blank Track in the sheet -> no forced human override
    assert rows[0]["screen_human"] == "needs_review"  # blank Screening -> safe default


def test_new_paper_fetches_immediate_citation_edges(tmp_path, config, conn):
    path = _make_sheet(tmp_path, [{"title": "", "arxiv_id": "2302.12173", "track": "Security", "screening": "Include"}])
    fake_paper = {
        "paper_id": "s2-abc123", "title": "T", "abstract": "A", "authors": [], "year": 2023,
        "venue": "V", "doi": None, "arxiv_id": "2302.12173", "url": None, "citation_count": 5,
    }
    fake_ref = {
        "paper_id": "s2-ref456", "title": "A Reference", "abstract": "", "authors": [], "year": 2023,
        "venue": "V", "doi": None, "arxiv_id": None, "url": None, "citation_count": 0,
    }
    with patch("src.clients.semantic_scholar.SemanticScholarClient.get_paper", return_value=fake_paper), \
         patch("src.clients.semantic_scholar.SemanticScholarClient.get_references", return_value={"papers": [fake_ref], "raw_pages": [], "hit_count": 1}), \
         patch("src.clients.semantic_scholar.SemanticScholarClient.get_citations", return_value={"papers": [], "raw_pages": [], "hit_count": 0}):
        counts = excel_sync.sync_excel(conn, config, path)

    assert counts["edges_added"] == 1
    edges = db.all_citation_edges(conn)
    assert len(edges) == 1
    assert edges[0]["citing_paper_id"] == "s2-abc123"
    assert edges[0]["cited_paper_id"] == "s2-ref456"
    assert db.get_paper(conn, "s2-ref456") is not None
