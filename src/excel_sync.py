"""Sync hand-added rows in the Excel data source into the SQLite DB.

A paper typed into the spreadsheet by hand only exists there until this
runs: it resolves full metadata via S2/arXiv/OpenAlex where possible (using
whatever arxiv_id/doi/title was given), inserts it as a real DB-backed paper
(discovered_via='manual'), and fetches its immediate references/citations so
it gets real citation-network edges instead of sitting as an isolated node.
The resolved paper_id and any backfilled metadata get written back into the
spreadsheet; a user's own Track/Screening choice for that row is never
touched - it becomes track_human/screen_human, same status as a seed's human
override (NFR-4: never silently overwritten by a later automated pass).

Rows already backed by a real DB paper_id are skipped (cheap to re-run after
adding a few more rows - already-synced ones are a no-op).
"""
from __future__ import annotations

import hashlib
import logging
from datetime import datetime, timezone

import openpyxl

from . import classify, db, dedup
from .clients.arxiv import ArxivClient, normalize_arxiv_id
from .clients.openalex import OpenAlexClient
from .clients.semantic_scholar import SemanticScholarClient
from .dedup import normalize_title
from .excel_source import COLUMNS, SHEET_NAME
from .labels import resolve_screen, resolve_track
from .snowball import cache_raw, record_edge, year_in_range
from .venues import normalize_venue

logger = logging.getLogger(__name__)


def _resolve_new_paper(row_vals: dict, s2, arxiv, openalex) -> dict:
    arxiv_id = normalize_arxiv_id(str(row_vals.get("arxiv_id") or ""))
    doi = (str(row_vals.get("doi") or "")).strip() or None
    resolved: dict = {}

    if arxiv_id:
        try:
            hit = s2.get_paper(f"ARXIV:{arxiv_id}")
            if hit:
                resolved = hit
        except Exception as exc:
            logger.warning("S2 lookup failed for arXiv:%s (%s)", arxiv_id, exc)
        if not resolved.get("abstract"):
            try:
                hit = arxiv.get_by_id(arxiv_id)
                if hit:
                    resolved = dedup.merge_fill_gaps(resolved, hit) if resolved else hit
            except Exception as exc:
                logger.warning("arXiv lookup failed for %s (%s)", arxiv_id, exc)
    elif doi:
        try:
            hit = s2.get_paper(f"DOI:{doi}")
            if hit:
                resolved = hit
        except Exception as exc:
            logger.warning("S2 lookup failed for DOI:%s (%s)", doi, exc)
        if not resolved.get("abstract"):
            try:
                hit = openalex.get_by_doi(doi)
                if hit:
                    resolved = dedup.merge_fill_gaps(resolved, hit) if resolved else hit
            except Exception as exc:
                logger.warning("OpenAlex lookup failed for DOI:%s (%s)", doi, exc)
    elif row_vals.get("title"):
        try:
            result = s2.search(str(row_vals["title"]), year_start=2000, max_results=1)
            if result["papers"]:
                resolved = result["papers"][0]
        except Exception as exc:
            logger.warning("S2 title search failed for %r (%s)", row_vals["title"], exc)

    if not resolved.get("arxiv_id") and arxiv_id:
        resolved["arxiv_id"] = arxiv_id
    if not resolved.get("doi") and doi:
        resolved["doi"] = doi
    return resolved


def _fetch_immediate_edges(conn, config, s2, paper_id: str, raw_cache_dir) -> dict:
    """One hop of backward+forward from this single new paper - not a full
    recursive snowball, just enough for it to show real connections."""
    now = datetime.now(timezone.utc).isoformat()
    counts = {"edges": 0, "new_related": 0}
    for direction, endpoint in (("backward", "get_references"), ("forward", "get_citations")):
        try:
            result = getattr(s2, endpoint)(paper_id)
        except Exception as exc:
            logger.warning("S2 %s lookup failed for %s: %s", endpoint, paper_id, exc)
            continue
        cache_raw(raw_cache_dir, f"manual_sync_{direction}", result["raw_pages"])
        for related in result["papers"]:
            related_id = related.get("paper_id")
            if not related_id:
                continue
            existing_id = dedup.resolve_existing_paper_id(conn, related)
            final_id = existing_id or related_id
            record_edge(conn, paper_id, final_id, direction)
            counts["edges"] += 1
            if existing_id or not year_in_range(related.get("year"), config):
                continue
            tags = classify.classify_paper(related.get("title"), related.get("abstract"), config)
            record = {
                **related, "paper_id": final_id, "doi": dedup.normalize_doi(related.get("doi")),
                "track_auto": tags["track_auto"], "screen_auto": tags["screen_auto"],
                "discovered_via": direction, "hop": 1, "source_paper_id": paper_id,
                "discovered_query": None, "seed_category": None, "date_added": now,
            }
            db.upsert_paper(conn, record)
            counts["new_related"] += 1
    return counts


def sync_excel(conn, config, excel_path) -> dict:
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    col_idx = {name: header.index(name) + 1 for name in COLUMNS if name in header}

    s2 = SemanticScholarClient(config)
    arxiv = ArxivClient(config)
    openalex = OpenAlexClient(config)
    raw_cache_dir = config.path("raw_cache_dir")
    now = datetime.now(timezone.utc).isoformat()

    counts = {"scanned": 0, "already_synced": 0, "newly_synced": 0, "edges_added": 0}

    for row_idx in range(2, ws.max_row + 1):
        row_vals = {name: ws.cell(row=row_idx, column=idx).value for name, idx in col_idx.items()}
        # a row is "worth processing" if it has a title OR an identifier -
        # someone might paste just an arxiv_id/doi and let sync fill in the rest
        if not (row_vals.get("title") or row_vals.get("arxiv_id") or row_vals.get("doi")):
            continue
        counts["scanned"] += 1

        existing_pid = str(row_vals.get("paper_id") or "").strip()
        if existing_pid and not existing_pid.startswith("manual:") and db.get_paper(conn, existing_pid):
            counts["already_synced"] += 1
            continue
        resolved = _resolve_new_paper(row_vals, s2, arxiv, openalex)

        candidate = dedup.merge_fill_gaps(resolved, {
            "title": row_vals.get("title"),
            "authors": [a.strip() for a in str(row_vals.get("authors") or "").split(";") if a.strip()],
            "year": row_vals.get("year"),
            "venue": row_vals.get("venue"),
            "doi": row_vals.get("doi"),
            "arxiv_id": normalize_arxiv_id(str(row_vals.get("arxiv_id") or "")) or None,
            "url": row_vals.get("url"),
        })

        existing_id = dedup.resolve_existing_paper_id(conn, candidate)
        paper_id = (
            existing_id
            or resolved.get("paper_id")
            or (f"arxiv:{candidate['arxiv_id']}" if candidate.get("arxiv_id") else None)
            or ("manual:" + hashlib.sha1(normalize_title(candidate.get("title") or "").encode()).hexdigest()[:16])
        )

        if not existing_id:
            track_code = resolve_track(row_vals.get("track"))
            track_human = track_code if track_code != "Unclear" else None
            screen_human = resolve_screen(row_vals.get("screening"))
            tags = classify.classify_paper(candidate.get("title"), candidate.get("abstract"), config)
            record = {
                **candidate,
                "paper_id": paper_id,
                "doi": dedup.normalize_doi(candidate.get("doi")),
                "track_auto": tags["track_auto"],
                "track_human": track_human,
                "screen_auto": tags["screen_auto"],
                "screen_human": screen_human,
                "discovered_via": "manual",
                "hop": 0,
                "source_paper_id": None,
                "discovered_query": None,
                "seed_category": None,
                "date_added": now,
            }
            db.upsert_paper(conn, record)
            conn.commit()

            if not paper_id.startswith("manual:") and not paper_id.startswith("arxiv:"):
                edge_counts = _fetch_immediate_edges(conn, config, s2, paper_id, raw_cache_dir)
                counts["edges_added"] += edge_counts["edges"]
                conn.commit()

        counts["newly_synced"] += 1
        ws.cell(row=row_idx, column=col_idx["paper_id"], value=paper_id)
        final_row = db.get_paper(conn, paper_id)
        if final_row:
            if not row_vals.get("abstract") and final_row["abstract"]:
                ws.cell(row=row_idx, column=col_idx["abstract"], value=final_row["abstract"])
            if not row_vals.get("citation_count") and final_row["citation_count"]:
                ws.cell(row=row_idx, column=col_idx["citation_count"], value=final_row["citation_count"])
            if not row_vals.get("venue") and final_row["venue"]:
                ws.cell(row=row_idx, column=col_idx["venue"], value=normalize_venue(final_row["venue"]))
            if not row_vals.get("year") and final_row["year"]:
                ws.cell(row=row_idx, column=col_idx["year"], value=final_row["year"])

    wb.save(excel_path)
    return counts
