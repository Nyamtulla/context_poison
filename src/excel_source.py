"""The editable Excel workbook that acts as the dashboard's data source.

The SQLite DB remains the pipeline's actual source of truth (search,
snowball, replicability all depend on it) - this is a one-way bootstrap
from DB -> Excel, created once (by default from the high-confidence pool in
src/confidence.py). After that, the spreadsheet is what the dashboard reads
on every load: add a row for a paper you found by hand, change a Track or
Screening decision, edit any metadata - refresh the dashboard and it's
there. Nothing here writes back into the DB; hand edits live in the
spreadsheet only, which is exactly what makes this fast to use without
touching SQL or the CLI.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import db
from .dedup import normalize_title
from .labels import SCREEN_DISPLAY_NAMES, TRACK_DISPLAY_NAMES, resolve_screen, resolve_track
from .venues import normalize_venue

SHEET_NAME = "Papers"
COLUMNS = [
    "paper_id", "title", "authors", "year", "venue", "track", "screening",
    "citation_count", "doi", "arxiv_id", "url", "pdf_local_path",
    "discovered_via", "discovered_query", "hop", "seed_category", "abstract",
]
_COLUMN_WIDTHS = {
    "title": 50, "authors": 25, "abstract": 60, "venue": 18,
    "url": 30, "pdf_local_path": 30, "doi": 20, "arxiv_id": 12,
}


def _row_from_db_row(row) -> dict:
    authors = row["authors"]
    if authors:
        try:
            authors = "; ".join(json.loads(authors))
        except (TypeError, ValueError):
            authors = ""
    else:
        authors = ""
    track_code = row["track_human"] or row["track_auto"] or "Unclear"
    screen_code = row["screen_human"] or row["screen_auto"] or "needs_review"
    return {
        "paper_id": row["paper_id"],
        "title": row["title"] or "",
        "authors": authors,
        "year": row["year"],
        "venue": normalize_venue(row["venue"]) or "",
        "track": TRACK_DISPLAY_NAMES.get(track_code, track_code),
        "screening": SCREEN_DISPLAY_NAMES.get(screen_code, screen_code),
        "citation_count": row["citation_count"] or 0,
        "doi": row["doi"] or "",
        "arxiv_id": row["arxiv_id"] or "",
        "url": row["url"] or "",
        "pdf_local_path": row["pdf_local_path"] or "",
        "discovered_via": row["discovered_via"] or "",
        "discovered_query": row["discovered_query"] or "",
        "hop": row["hop"],
        "seed_category": row["seed_category"] or "",
        "abstract": row["abstract"] or "",
    }


def build_workbook(conn, ids: set[str] | None = None) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for row in db.all_papers(conn):
        if ids is not None and row["paper_id"] not in ids:
            continue
        r = _row_from_db_row(row)
        ws.append([r[c] for c in COLUMNS])

    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _COLUMN_WIDTHS.get(col, 14)
    return wb


def export_to_path(conn, path: str | Path, ids: set[str] | None = None, overwrite: bool = False) -> Path:
    path = Path(path)
    if path.exists() and not overwrite:
        raise FileExistsError(
            f"{path} already exists - refusing to overwrite hand-edited data without overwrite=True"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(conn, ids=ids).save(path)
    return path


def load_dataframe(path: str | Path) -> pd.DataFrame:
    """Read the spreadsheet back into the same shape the dashboard expects
    (effective_track/effective_screen/venue_short/authors-as-list/etc) -
    tolerant of hand edits: blank paper_id (a manually added row) gets a
    stable ID derived from its title; unrecognized Track/Screening text
    falls back to Unclear/Needs Review rather than breaking silently."""
    df = pd.read_excel(path, sheet_name=SHEET_NAME, engine="openpyxl")
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = None

    df["authors"] = df["authors"].fillna("").apply(
        lambda s: [a.strip() for a in str(s).split(";") if a.strip()]
    )
    df["year"] = pd.to_numeric(df["year"], errors="coerce")
    df["citation_count"] = pd.to_numeric(df["citation_count"], errors="coerce").fillna(0)
    df["hop"] = pd.to_numeric(df["hop"], errors="coerce").fillna(0).astype(int)
    df["effective_track"] = df["track"].apply(resolve_track)
    df["effective_screen"] = df["screening"].apply(resolve_screen)
    df["venue_short"] = df["venue"].apply(normalize_venue)

    missing_id = df["paper_id"].isna() | (df["paper_id"].astype(str).str.strip().isin(["", "nan"]))
    if missing_id.any():
        df.loc[missing_id, "paper_id"] = df.loc[missing_id, "title"].apply(
            lambda t: "manual:" + hashlib.sha1(normalize_title(str(t)).encode()).hexdigest()[:16]
        )
    return df
