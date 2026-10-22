"""Write the full-text re-coded labels into the corpus spreadsheet.

The re-coding produced CSVs; the registries were updated from them by
propagate_recode.py. But the taxonomy cube in gen_paper_numbers.py reads
channel/consequence/track from data/exports/paper_dashboard_source.xlsx, so
without this step the cube still reports the title+abstract labels.

A `none` is written through as "none": it does not match the cube's channel
or consequence list, so the paper drops out of the cube, which is the
intended behaviour for a survey spanning several routes co-equally.

Backs the workbook up first, following the existing backup-<stamp> naming.
"""
from __future__ import annotations

import csv
import glob
import pathlib
import shutil
import time

import openpyxl

REPO = pathlib.Path(__file__).resolve().parent.parent
XLSX = REPO / "data" / "exports" / "paper_dashboard_source.xlsx"
RAW = REPO / "data" / "registries" / "raw"
AXES = ("track", "channel", "consequence")


def main():
    new = {}
    for f in sorted(glob.glob(str(RAW / "recode_decisions_batch*.csv"))):
        for r in csv.DictReader(open(f, encoding="utf-8")):
            new[r["paper_id"]] = {a: (r[a] or "").strip() for a in AXES}
    print(f"re-coded papers: {len(new)}")

    stamp = time.strftime("%Y%m%d-%H%M%S")
    shutil.copy(XLSX, XLSX.with_name(f"paper_dashboard_source.backup-recode-{stamp}.xlsx"))

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Papers"]
    hdr = [c.value for c in ws[1]]
    col = {a: hdr.index(a) + 1 for a in AXES}
    pid = hdr.index("paper_id") + 1

    changed = {a: 0 for a in AXES}
    touched = 0
    for row in range(2, ws.max_row + 1):
        p = str(ws.cell(row=row, column=pid).value)
        if p not in new:
            continue
        hit = False
        for a in AXES:
            cur = str(ws.cell(row=row, column=col[a]).value or "").strip()
            n = new[p][a]
            if n and n != cur:
                ws.cell(row=row, column=col[a]).value = n
                changed[a] += 1
                hit = True
        touched += hit
    wb.save(XLSX)
    print(f"rows touched: {touched}")
    for a in AXES:
        print(f"  {a:12s} updated {changed[a]}")
    print(f"backup: paper_dashboard_source.backup-recode-{stamp}.xlsx")


if __name__ == "__main__":
    main()
