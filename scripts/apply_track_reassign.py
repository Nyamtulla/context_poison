"""Retire the `Both` track: assign every paper to Security or ML/AI.

The cube's third axis is intent, but the corpus has no intent column, so
track stood in for it and `Both` was filtered out of both fig_cube and the
empty-cell macro. That left 63 of 189 cells unoccupiable by construction and
inflated the published empty-cell rate.

Decisions come from data/registries/raw/track_reassign_both.csv, coded under
the rule pre-registered in track_reassignment_prediction.md. Writes the
paper rows AND the registry entries, because build_timeline.py prefers the
registry's own track over the paper's.
"""
import csv, json, pathlib, shutil, datetime, collections
import openpyxl

REPO = pathlib.Path(__file__).resolve().parent.parent
XL = REPO / "data/exports/paper_dashboard_source.xlsx"
REG = REPO / "data/registries"
DEC = REG / "raw/track_reassign_both.csv"

dec = {r["paper_id"]: r["new_track"] for r in csv.DictReader(DEC.open())}
print(f"{len(dec)} decisions, " + ", ".join(
    f"{k}={v}" for k, v in collections.Counter(dec.values()).items()))

stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
shutil.copy(XL, XL.with_suffix(f".backup-track-{stamp}.xlsx"))

def match(pid):
    """Decisions carry truncated ids; the corpus carries full ones."""
    if pid in dec:
        return dec[pid]
    for k, v in dec.items():
        if pid and str(pid).startswith(k):
            return v
    return None

wb = openpyxl.load_workbook(XL)
ws = wb["Papers"]
hdr = [c.value for c in ws[1]]
ti, pi = hdr.index("track") + 1, hdr.index("paper_id") + 1
n = 0
for row in range(2, ws.max_row + 1):
    if str(ws.cell(row, ti).value or "").strip() != "Both":
        continue
    new = match(ws.cell(row, pi).value)
    if new:
        ws.cell(row, ti).value = new
        n += 1
    else:
        print(f"  UNMATCHED paper row {row}: {ws.cell(row, pi).value}")
wb.save(XL)
print(f"papers: {n} rows retracked")

for fn in ("rq3_pollution_registry.json", "rq4_defense_registry.json"):
    p = REG / fn
    rows = json.loads(p.read_text())
    m = 0
    for e in rows:
        if (e.get("track") or "").strip() != "Both":
            continue
        new = match(e.get("paper_id"))
        if new:
            e["track"] = new
            m += 1
        else:
            print(f"  UNMATCHED {fn}: {e.get('paper_id')} {e.get('title','')[:50]}")
    p.write_text(json.dumps(rows, indent=1))
    print(f"{fn}: {m} entries retracked")
