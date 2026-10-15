"""Build the contribution re-screen: what does each paper actually contribute?

Every number in this SoK rests on a census of named attacks, mechanisms,
defenses and fixes. Two failures show the census cannot be trusted as built.
RhinoInsight was registered as a fix for context rot on the strength of its
framing, while its evaluation is a report-quality score. `is_extension_of` is
populated on 5% of defense entries, so the registry cannot tell a method a
paper invented from one it applied off the shelf.

This pass asks one question per paper, answerable from the abstract, the
contributions list and the evaluation section, and refuses to be answered from
the introduction alone.

WHAT THIS IS NOT. It is not "delete every paper with no named technique". 302
of the 982 papers name neither a mechanism nor a defense, and they include
AgentDojo, WebArena, InjecAgent, RULER, Agent Security Bench and the Knowledge
Conflicts survey. AgentDojo is the harness 56 of this project's own
experiments run on and is cited in the paper. Benchmarks, surveys and
measurement studies belong in the corpus: they are how the field evaluates,
and they carry most of the cross-track citation bridges. Only `out_of_scope`
removes a paper.

    python3 scripts/build_contribution_screen.py            # all included papers
    python3 scripts/build_contribution_screen.py --per 60   # packet size
"""
import argparse
import collections
import json
import pathlib

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = pathlib.Path(__file__).resolve().parent.parent
R = REPO / "data" / "registries"
OUT = REPO / "data" / "exports" / "contribution_screen"

CONTRIB = [
 ("proposes_attack", "The paper introduces a NEW named attack technique.",
  "keep. It is a Track A mechanism."),
 ("proposes_mechanism", "Introduces a NEW named non-adversarial failure mode "
  "(context rot, lost-in-the-middle and the like).", "keep. Track B mechanism."),
 ("proposes_defense", "Introduces a NEW named defense against a deliberate attack.",
  "keep. Track A defense."),
 ("proposes_fix", "Introduces a NEW named method for a non-adversarial failure.",
  "keep. Track B fix."),
 ("improves_existing", "Extends or modifies a technique that already existed. "
  "Name what it extends.", "keep, and record the parent."),
 ("applies_existing", "Uses an existing technique without changing it. Evaluates "
  "with it, builds on it, reports it as a baseline.",
  "keep as a paper, but it must NOT hold a registry entry of its own."),
 ("benchmark_or_dataset", "Contributes an evaluation artifact rather than a "
  "technique: AgentDojo, InjecAgent, RULER, WebArena.",
  "KEEP. These are how the field evaluates and they carry the cross-track "
  "citation bridges."),
 ("measurement_study", "Measures or characterises a phenomenon without "
  "proposing a technique.", "KEEP. This is evidence, not absence."),
 ("survey_or_sok", "Surveys or systematizes existing work.",
  "KEEP, and it is also a competitor to compare against."),
 ("position_or_conceptual", "Argues a position, proposes an architecture with "
  "no implementation or evaluation.",
  "keep, but flag: it should not count as a tested defense."),
 ("out_of_scope", "Does not study context becoming untrustworthy at all. It "
  "studies agent capability, planning or task performance and uses the words "
  "'context' or 'long-horizon' in another sense.",
  "THE ONLY VALUE THAT REMOVES A PAPER. RhinoInsight and FS-Researcher were "
  "removed under this value on 2026-10-15."),
]

FIELDS = [
 ("contribution", "pick one", "The single best fit from the Contribution sheet. "
  "If two fit, pick the one the paper's own contributions list leads with."),
 ("technique_name", "free text", "The name the paper gives its technique, exactly "
  "as the paper writes it. Blank for benchmark, survey, measurement, "
  "out_of_scope."),
 ("origin", "invents / improves / applies / none",
  "invents = the paper is the first to propose it. improves = it changes "
  "something that existed. applies = it uses someone else's unchanged. "
  "none = no technique."),
 ("parent_technique", "free text", "If origin is improves or applies, what does "
  "it build on? Blank otherwise."),
 ("evaluated_against", "free text",
  "What the EVALUATION SECTION actually runs the method against. Attack names, "
  "benchmark names, 'no evaluation'. NOT what the introduction claims to "
  "address. This is the field that catches a paper claiming one thing and "
  "measuring another."),
 ("evidence_quote", "free text",
  "One verbatim sentence from the paper supporting `contribution` and `origin`, "
  "with its section. NO QUOTE, NO ENTRY."),
 ("names_registry_technique", "yes / no / unsure",
  "THE AGENTDOJO QUESTION. Does this paper name a specific attack or failure "
  "mode that other papers then cite and test against, even if the paper is a "
  "benchmark, a harness or a survey? AgentDojo is a harness AND it names seven "
  "injections that 38 confirmed pairs depend on, so the answer for it is yes. "
  "WebArena is a harness that names none, so no. Answer from the paper's own "
  "section headings and tables, not from its genre."),
 ("named_techniques", "free text",
  "If names_registry_technique is yes, list them as the paper names them, "
  "semicolon separated. 'important-instructions injection; ignore-previous "
  "injection'. Blank otherwise."),
 ("belongs_in_corpus", "yes / no / unsure",
  "no ONLY for out_of_scope. A benchmark, a survey and a measurement study all "
  "belong. If unsure, say unsure and explain in notes."),
 ("notes", "free text", "Anything that does not fit above, especially "
  "disagreement with the current coding."),
]


def header(ws, cols, width=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2A78D6")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[chr(64+i)].width = (width or [18]*len(cols))[i-1]
    ws.freeze_panes = "A2"


def build(no, papers, total_packets):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "start here"
    ws.column_dimensions["A"].width = 118
    lines = [
     (f"Contribution re-screen, packet {no} of {total_packets}: "
      f"{len(papers)} papers", True),
     ("", False),
     ("ONE QUESTION: what does this paper actually contribute, and did it "
      "invent that, improve it, or just use it?", True),
     ("", False),
     ("Why this exists. Every count in the study rests on a census of named "
      "techniques. Two papers were found registered as fixes for context rot "
      "on the strength of their framing alone; their evaluations measured "
      "report quality. We need to know how often that happened.", False),
     ("", False),
     ("WHAT NOT TO DO. Do not mark a paper out_of_scope because it proposes no "
      "technique. Benchmarks, surveys and measurement studies all belong in "
      "the corpus. out_of_scope means the paper is not about untrustworthy "
      "context at all.", True),
     ("", False),
     ("THE COLUMN THAT MATTERS MOST IS names_registry_technique. A paper can be "
      "a harness and still name attacks. AgentDojo is an evaluation environment "
      "AND it names seven injections; 38 of the 349 confirmed "
      "(defense, attack) pairs in this study depend on them, so it is a "
      "contributing paper by any rule. WebArena is also a harness and names "
      "none. Genre does not decide this. Read the paper's section headings and "
      "tables and answer what it NAMES.", True),
     ("", False),
     ("Where to look, in this order:", True),
     ("   1. the abstract and the contributions list  -> contribution, origin", False),
     ("   2. the EVALUATION or EXPERIMENTS section    -> evaluated_against", False),
     ("   3. related work                             -> parent_technique", False),
     ("", False),
     ("Do not answer from the introduction alone. An introduction saying 'we "
      "address context rot' is not evidence that context rot was measured. "
      "That specific confusion is what this pass exists to find.", False),
     ("", False),
     ("Every row needs a verbatim quote. No quote, no entry.", True),
     ("", False),
     ("Sheets: 'contribution' = the eleven values and what each one means. "
      "'fields' = what each column wants. 'papers' = your list, one row each.", False),
    ]
    for txt, bold in lines:
        c = ws.cell(ws.max_row + 1, 1, txt)
        c.font = Font(bold=bold, size=12 if bold and txt else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wsc = wb.create_sheet("contribution")
    header(wsc, ["value", "means", "what happens to the paper"], [24, 62, 54])
    for v, m, w in CONTRIB:
        r = wsc.max_row + 1
        wsc.cell(r, 1, v); wsc.cell(r, 2, m); wsc.cell(r, 3, w)
        for cc in range(1, 4):
            wsc.cell(r, cc).alignment = Alignment(wrap_text=True, vertical="top")

    wsf = wb.create_sheet("fields")
    header(wsf, ["column", "allowed values", "what it wants"], [22, 30, 72])
    for f, a, d in FIELDS:
        r = wsf.max_row + 1
        wsf.cell(r, 1, f); wsf.cell(r, 2, a); wsf.cell(r, 3, d)
        for cc in range(1, 4):
            wsf.cell(r, cc).alignment = Alignment(wrap_text=True, vertical="top")

    ws2 = wb.create_sheet("papers")
    cols = ["#", "title", "year", "venue", "track", "pdf",
            "current registry entry", "current origin (auto)"] + [f for f, _, _ in FIELDS]
    header(ws2, cols, [5, 56, 7, 16, 10, 40, 40, 18]
           + [22, 26, 16, 22, 34, 44, 24, 34, 16, 30])
    for i, p in enumerate(papers, 1):
        r = ws2.max_row + 1
        for j, v in enumerate([i, p["title"], p["year"], p["venue"], p["track"],
                               p["pdf"], p["registry"], p["auto_origin"]], 1):
            ws2.cell(r, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / f"contribution-screen-{no:02d}.xlsx")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--per", type=int, default=60)
    a = ap.parse_args()

    wb = openpyxl.load_workbook(
        REPO / "data" / "exports" / "paper_dashboard_source.xlsx", read_only=True)
    ws = wb["Papers"]; rows = list(ws.iter_rows(values_only=True)); wb.close()
    hdr = list(rows[0])
    I = {k: hdr.index(k) for k in
         ("paper_id", "title", "year", "venue", "track", "screening",
          "pdf_local_path")}
    inc = [r for r in rows[1:] if str(r[I["screening"]]).strip() != "Exclude"]

    reg = collections.defaultdict(list)
    for fn, kind in (("rq3_pollution_registry.json", "mechanism"),
                     ("rq4_defense_registry.json", "defense")):
        for e in json.loads((R / fn).read_text()):
            reg[e.get("paper_id")].append(f"{kind}: {e.get('name','')[:46]}")
    auto = {}
    orig = R / "defense_origin.json"
    if orig.exists():
        for d in json.loads(orig.read_text())["defenses"]:
            auto[d["paper_id"]] = d["origin"]

    rank = {p["paper_id"]: p["score"] for p in
            json.loads((R / "paper_importance_ranking.json").read_text())["papers"]}
    papers = [{"paper_id": r[I["paper_id"]], "title": r[I["title"]],
               "year": r[I["year"]], "venue": r[I["venue"]], "track": r[I["track"]],
               "pdf": r[I["pdf_local_path"]],
               "registry": "; ".join(reg.get(r[I["paper_id"]], [])) or "(none)",
               "auto_origin": auto.get(r[I["paper_id"]], "")}
              for r in inc]
    # most load-bearing first, so an unfinished pass still covers what matters
    papers.sort(key=lambda p: -rank.get(p["paper_id"], 0))

    n = (len(papers) + a.per - 1) // a.per
    for i in range(n):
        build(i + 1, papers[i*a.per:(i+1)*a.per], n)
    print(f"{len(papers)} papers -> {n} packets of up to {a.per}, in {OUT}")
    print("ordered by importance score, so packet 01 holds the most load-bearing papers")


if __name__ == "__main__":
    main()
