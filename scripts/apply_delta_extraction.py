"""Fold the screening-delta extraction back into the corpus.

Reads the subagent outputs (data/registries/raw/delta_extract_out*.jsonl) plus
the earlier hand-coded batch (screening_delta_coding_batch1.json), then:

  1. writes the 15 full-text analysis columns into paper_dashboard_source.xlsx
     for the delta rows (the Excel is the ONLY store of these columns -- the
     SQLite DB does not carry them, so this file can never be regenerated from
     the DB and is only ever appended to / filled in, never rebuilt);
  2. emits RQ3/RQ4 raw batch CSVs in the exact shape build_registry.py reads.

    python3 scripts/apply_delta_extraction.py [--dry-run]
"""
from __future__ import annotations
import argparse, csv, json, sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
RAW = REPO / "data/registries/raw"
XL = REPO / "data/exports/paper_dashboard_source.xlsx"

ANALYSIS = ["channel","consequence","defense_intervention_point","temporal_persistence",
            "cites_track_a","cites_track_b","evidence_grade","artifacts_released",
            "models_evaluated","datasets_benchmarks","key_result","baselines_compared",
            "threat_model","stated_limitations","technical_summary"]

CHANNELS = {"tool-output","direct-input","RAG","memory","tool-metadata","cross-modal",
            "multi-agent","skill","supply-chain"}
CONSEQ = {"goal-hijack","reasoning-corruption","silent-corruption","persistence-backdoor",
          "data-exfiltration","resource-abuse"}
STAGES = {"ingestion","reasoning","execution","none"}


def _packet_owner():
    """row -> the packet number whose packet file actually contains that row.

    A couple of agents in the first (rate-limited) run read a sibling's packet
    and wrote a handful of rows into the wrong output file. The copies agree on
    every controlled-vocabulary field and differ only in prose phrasing, but
    "whichever file sorted last" is not a defensible tie-break for research
    data. Resolve by ownership instead, so the result is deterministic and
    re-running cannot silently change which extraction won."""
    owner = {}
    for pf in RAW.glob("delta_extract_packet*.json"):
        n = int("".join(c for c in pf.stem if c.isdigit()))
        for entry in json.loads(pf.read_text()):
            owner[int(entry["row"])] = n
    return owner


def load_jsonl():
    out, bad, dupes = {}, [], []
    owner = _packet_owner()
    for fp in sorted(RAW.glob("delta_extract_out*.jsonl")):
        fnum = int("".join(c for c in fp.stem if c.isdigit()))
        for ln, line in enumerate(fp.read_text().splitlines(), 1):
            line = line.strip()
            if not line:
                continue
            try:
                r = json.loads(line)
            except json.JSONDecodeError as e:
                bad.append(f"{fp.name}:{ln} {e}")
                continue
            if "row" not in r:
                bad.append(f"{fp.name}:{ln} no row")
                continue
            row = int(r["row"])
            r["_from_packet"] = fnum
            if row in out:
                prev = out[row]["_from_packet"]
                keep_new = (owner.get(row) == fnum) and prev != owner.get(row)
                dupes.append(f"row {row}: in out{prev} and out{fnum}; "
                             f"owner=packet{owner.get(row)} -> kept "
                             f"out{fnum if keep_new else prev}")
                if not keep_new:
                    continue
            out[row] = r
    return out, bad, dupes


def merge_batch1_coding(recs):
    """Attach the 2026-09-13 hand-coded RQ3/RQ4 judgments to their Excel rows.

    Those 40 papers were coded before the full-text extraction pass, into
    screening_delta_coding_batch1.json keyed by a 10-char paper_id prefix. The
    extraction agents were told needs_rq34=false for them, so their rq3/rq4
    blocks live only in that file and must be joined back in here or the
    registry silently loses 40 papers' worth of judgment.

    rq4 entries there carry name+stage only; channel/consequence/
    validated_against were backfilled separately into
    batch1_rq4_backfill.jsonl and are merged on top."""
    b1_path = REPO / "data/registries/screening_delta_coding_batch1.json"
    rows_path = REPO / "data/registries/screening_delta_excel_rows.json"
    if not (b1_path.exists() and rows_path.exists()):
        return 0, 0, []
    coded = json.loads(b1_path.read_text())["coded"]
    row_by10 = {r["paper_id"][:10]: int(r["row"])
                for r in json.loads(rows_path.read_text())}

    backfill = {}
    bf = RAW / "batch1_rq4_backfill.jsonl"
    if bf.exists():
        for line in bf.read_text().splitlines():
            if line.strip():
                b = json.loads(line)
                backfill[int(b["row"])] = b

    n3 = n4 = 0
    orphans = []
    for c in coded:
        row = row_by10.get(c["paper_id"])
        if row is None:
            orphans.append(c.get("title", c["paper_id"]))
            continue
        if row not in recs:
            orphans.append(f"row {row} coded but not extracted: {c.get('title','')[:50]}")
            continue
        r = recs[row]
        if isinstance(c.get("rq3"), dict):
            r["rq3"] = c["rq3"]
            n3 += int(bool(c["rq3"].get("has")))
        if isinstance(c.get("rq4"), dict):
            r4 = dict(c["rq4"])
            if r4.get("has") and row in backfill:
                b = backfill[row]
                for k in ("channel", "consequence", "validated_against"):
                    if b.get(k):
                        r4[k] = b[k]
                if b.get("conf"):
                    r4["conf"] = b["conf"]
            r["rq4"] = r4
            n4 += int(bool(r4.get("has")))
        if c.get("scope_verdict"):
            r["scope_verdict"] = c["scope_verdict"]
    return n3, n4, orphans


def repair_rq4_stage(recs):
    """Fill a missing rq4.stage from the paper-level defense_intervention_point.

    Agents occasionally emit an rq4 block without `stage`. The paper-level
    `defense_intervention_point` is the same judgment about the same defense, so
    when it is present and not "none" it is a sound, deterministic fallback.
    `validated_against` is NOT inferred -- that one depends on what the paper's
    evaluation actually ran, which cannot be recovered from other fields, so a
    missing value stays a hard validation failure to be resolved explicitly."""
    fixed = []
    for row, r in sorted(recs.items()):
        s4 = r.get("rq4")
        if isinstance(s4, dict) and s4.get("has") and not s4.get("stage"):
            dip = r.get("defense_intervention_point")
            if dip and dip != "none":
                s4["stage"] = dip
                fixed.append((row, dip))
    return fixed


def validate(recs):
    """Closed-vocabulary check. Returns per-row list of problems."""
    problems = []
    for row, r in sorted(recs.items()):
        if r.get("channel") not in CHANNELS:
            problems.append((row, "channel", r.get("channel")))
        if r.get("consequence") not in CONSEQ:
            problems.append((row, "consequence", r.get("consequence")))
        if r.get("defense_intervention_point") not in STAGES:
            problems.append((row, "defense_intervention_point", r.get("defense_intervention_point")))
        for k in ("cites_track_a","cites_track_b"):
            if str(r.get(k)).upper() not in {"Y","N"}:
                problems.append((row, k, r.get(k)))
        if str(r.get("evidence_grade")).upper() not in {"A","B","C","D"}:
            problems.append((row, "evidence_grade", r.get("evidence_grade")))
        for sub in ("rq3","rq4"):
            s = r.get(sub)
            if isinstance(s, dict) and s.get("has"):
                if not (s.get("name") or "").strip():
                    problems.append((row, f"{sub}.name", "empty"))
                if s.get("channel") and s["channel"] not in CHANNELS:
                    problems.append((row, f"{sub}.channel", s["channel"]))
                if s.get("consequence") and s["consequence"] not in CONSEQ:
                    problems.append((row, f"{sub}.consequence", s["consequence"]))
                if sub == "rq4":
                    if s.get("stage") not in STAGES - {"none"}:
                        problems.append((row, "rq4.stage", s.get("stage")))
                    if s.get("validated_against") not in {"adversarial","incidental","both"}:
                        problems.append((row, "rq4.validated_against", s.get("validated_against")))
    return problems


def normalize_evidence_grade(recs, xl_meta):
    """Enforce the corpus's venue-based evidence_grade convention.

    `evidence_grade` grades PUBLICATION RIGOR, not evaluation quality
    (context_integrity_sok_project_plan.md: peer-reviewed+artifacts / peer-reviewed
    / credible preprint / gray literature). The extraction spec originally stated
    it the other way round, so the subagents graded evaluation strength and drifted
    badly off the corpus convention.

    Two rules the existing 1,008-paper corpus never once violates, applied here
    deterministically rather than by re-running judgment:

      1. An arXiv preprint is C. 707 of 707 arXiv papers in the corpus are C;
         not one is B. A B on an arXiv row is always wrong.
      2. D means NO VENUE (thesis, unpublished). Both of the corpus's only two
         D papers have venue=None. A weak paper at a real venue is C, never D.

    The B-vs-C split *within* real venues stays as the agent judged it -- that one
    is a genuine venue-tier call, and the delta's rate (63.6%) sits close enough to
    the corpus's (48.8%) to be credible for a cohort that is, by construction, the
    paywalled peer-reviewed tail.
    """
    changed = []
    for row, r in sorted(recs.items()):
        venue, arxiv_id = xl_meta.get(row, ("", None))
        venue = (venue or "").strip()
        is_preprint = venue.lower() == "arxiv" or (not venue and arxiv_id)
        has_venue = bool(venue)
        g = str(r.get("evidence_grade") or "").upper()
        new = g
        if is_preprint and g == "B":
            new = "C"
        elif g == "D" and has_venue:
            new = "C"
        if new != g:
            r["evidence_grade"] = new
            changed.append((row, g, new, venue or "(none)"))
    return changed


def write_excel(recs, dry):
    import openpyxl
    wb = openpyxl.load_workbook(XL)
    ws = wb["Papers"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: n + 1 for n, h in enumerate(hdr)}
    written = 0
    for row, r in sorted(recs.items()):
        pid_xl = ws.cell(row, idx["paper_id"]).value
        if pid_xl != r.get("paper_id"):
            print(f"  !! row {row}: paper_id mismatch excel={pid_xl} rec={r.get('paper_id')} -- SKIPPED")
            continue
        for col in ANALYSIS:
            v = r.get(col)
            if v is None:
                continue
            ws.cell(row, idx[col]).value = str(v).strip()
        written += 1
    if not dry:
        wb.save(XL)
    return written


def rq_csvs(recs, batch1_rows, dry):
    """Emit RQ3 (track A / track B) and RQ4 delta batch CSVs."""
    a, b, d4 = [], [], []
    for row, r in sorted(recs.items()):
        r3, r4 = r.get("rq3"), r.get("rq4")
        if r3 is None and r4 is None:
            continue
        track = r.get("_track") or "Security"
        pid, title = r.get("paper_id",""), r.get("title","")
        if isinstance(r3, dict):
            rec = {"row":row,"paper_id":pid,"title":title,
                   "channel":r3.get("channel",""),"consequence":r3.get("consequence",""),
                   "is_extension_of":r3.get("is_extension_of",""),
                   "confidence":r3.get("conf","") or r3.get("confidence",""),
                   "notes":(r3.get("note","") or "").replace("\n"," ")}
            if track == "ML/AI":
                b.append({**rec,"has_mechanism":"Y" if r3.get("has") else "N",
                          "mechanism_name":r3.get("name","") if r3.get("has") else ""})
            else:
                a.append({**rec,"has_technique":"Y" if r3.get("has") else "N",
                          "technique_name":r3.get("name","") if r3.get("has") else ""})
        if isinstance(r4, dict):
            d4.append({"row":row,"paper_id":pid,"title":title,
                       "has_defense":"Y" if r4.get("has") else "N",
                       "defense_name":r4.get("name","") if r4.get("has") else "",
                       "channel":r4.get("channel",""),"consequence":r4.get("consequence",""),
                       "defense_intervention_point":r4.get("stage","") if r4.get("has") else "",
                       "track":track,"validated_against":r4.get("validated_against",""),
                       "is_extension_of":r4.get("is_extension_of",""),
                       "confidence":r4.get("conf","") or r4.get("confidence",""),
                       "notes":(r4.get("note","") or "").replace("\n"," ")})
    specs = [
        ("rq3_delta_track_a.csv", a, ["row","paper_id","title","has_technique","technique_name",
                                      "channel","consequence","is_extension_of","confidence","notes"]),
        ("rq3_delta_track_b.csv", b, ["row","paper_id","title","has_mechanism","mechanism_name",
                                      "channel","consequence","is_extension_of","confidence","notes"]),
        ("rq4_delta_batch.csv",  d4, ["row","paper_id","title","has_defense","defense_name","channel",
                                      "consequence","defense_intervention_point","track",
                                      "validated_against","is_extension_of","confidence","notes"]),
    ]
    for fname, rows, cols in specs:
        print(f"  {fname}: {len(rows)} rows ({sum(1 for x in rows if x.get('has_technique')=='Y' or x.get('has_mechanism')=='Y' or x.get('has_defense')=='Y')} positive)")
        if dry:
            continue
        with open(RAW / fname, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            for x in rows:
                w.writerow({c: x.get(c, "") for c in cols})
    return len(a), len(b), len(d4)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    recs, bad, dupes = load_jsonl()
    print(f"loaded {len(recs)} extraction records; {len(bad)} unparseable; {len(dupes)} duplicate row(s) resolved by packet ownership")
    for x in dupes:
        print("   ", x)
    for x in bad[:10]:
        print("   ", x)

    m3, m4, orphans = merge_batch1_coding(recs)
    print(f"\nmerged 2026-09-13 batch1 coding: {m3} rq3 positives, {m4} rq4 positives")
    for o in orphans:
        print(f"   (orphan) {o}")

    staged = repair_rq4_stage(recs)
    if staged:
        print(f"\nrq4.stage filled from defense_intervention_point: {len(staged)}")
        for row, dip in staged:
            print(f"   row {row}: stage <- {dip}")

    probs = validate(recs)
    print(f"\nvocabulary problems: {len(probs)}")
    for row, field, val in probs[:25]:
        print(f"   row {row}: {field} = {val!r}")
    if probs and not args.dry_run:
        print("\nREFUSING to write with vocabulary problems outstanding. Fix or re-run those rows.")
        sys.exit(1)

    # stamp track from the Excel so RQ3 routes to the right track file
    import openpyxl
    wb = openpyxl.load_workbook(XL, read_only=True)
    ws = wb["Papers"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ti, pi = hdr.index("track"), hdr.index("paper_id")
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if n in recs:
            recs[n]["_track"] = row[ti]
    wb.close()

    xl_meta = {}
    wb2 = openpyxl.load_workbook(XL, read_only=True)
    ws2 = wb2["Papers"]
    h2 = [c.value for c in next(ws2.iter_rows(min_row=1, max_row=1))]
    vi, axi = h2.index("venue"), h2.index("arxiv_id")
    for n, row in enumerate(ws2.iter_rows(min_row=2, values_only=True), start=2):
        xl_meta[n] = (row[vi], row[axi])
    wb2.close()
    fixes = normalize_evidence_grade(recs, xl_meta)
    print(f"\nevidence_grade normalized to the corpus venue convention: {len(fixes)} corrected")
    for row, a, b, v in fixes[:40]:
        print(f"   row {row}: {a} -> {b}   venue={v[:34]}")

    print(f"\nwriting analysis columns ({'DRY RUN' if args.dry_run else 'live'}) ...")
    n = write_excel(recs, args.dry_run)
    print(f"  {n} rows updated")

    print("\nRQ3/RQ4 raw batches:")
    rq_csvs(recs, None, args.dry_run)
    print("\ndone." + ("  (dry run -- nothing written)" if args.dry_run else ""))


if __name__ == "__main__":
    main()
