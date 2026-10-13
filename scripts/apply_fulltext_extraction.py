"""Write the 30 late-arriving papers' full-text analysis columns into the Excel.

`data/exports/paper_dashboard_source.xlsx` is the only store of the analysis
columns - the SQLite DB is rebuilt from it - so the extraction lands here, not
in the database. The DB is then updated to match so the two do not drift.

These 30 arrived after the original extraction pass: 28 downloaded by hand for
the missing-PDF worklist and 2 fetched from arXiv. Until now their taxonomy
fields were coded from abstracts and their seven prose columns were empty.
"""
from __future__ import annotations

import glob
import json
import pathlib
import shutil
import sqlite3
import sys
import time

import openpyxl

REPO = pathlib.Path(__file__).resolve().parent.parent
XL = REPO / "data" / "exports" / "paper_dashboard_source.xlsx"
NEED = REPO / "verification" / "needs_fulltext_extraction.json"
FIELDS = ["models_evaluated", "datasets_benchmarks", "key_result",
          "baselines_compared", "threat_model", "stated_limitations",
          "technical_summary"]


def main() -> None:
    dry = "--apply" not in sys.argv
    papers = json.loads(NEED.read_text())["papers"]
    recs = []
    for f in sorted(glob.glob(sys.argv[1])):
        recs += json.loads(pathlib.Path(f).read_text())
    by_idx = {r["idx"]: r for r in recs}
    missing = [i for i in range(len(papers)) if i not in by_idx]
    if missing:
        raise SystemExit(f"no extraction for indices {missing}")
    pid_of = {i: papers[i]["paper_id"] for i in range(len(papers))}
    print(f"{len(by_idx)} extractions for {len(papers)} papers")

    wb = openpyxl.load_workbook(XL)
    ws = wb[wb.sheetnames[0]]
    col = {c.value: n + 1 for n, c in enumerate(ws[1])}
    for f in FIELDS + ["paper_id"]:
        if f not in col:
            raise SystemExit(f"workbook lacks column {f}")

    want = {pid_of[i]: by_idx[i] for i in by_idx}
    wrote = 0
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(row=r, column=col["paper_id"]).value
        rec = want.get(pid)
        if not rec:
            continue
        for f in FIELDS:
            ws.cell(row=r, column=col[f], value=rec[f])
        wrote += 1
    print(f"rows matched in the workbook: {wrote} / {len(want)}")
    if dry:
        print("\ndry run; pass --apply to write")
        return
    bak = XL.with_suffix(f".backup-{time.strftime('%Y%m%d-%H%M%S')}.xlsx")
    shutil.copy2(XL, bak)
    wb.save(XL)
    print(f"backed up -> {bak.name}\nwrote {XL.relative_to(REPO)}")

    pe = sqlite3.connect(REPO / "data" / "paper_explorer.db")
    for pid, rec in want.items():
        pe.execute(f"update paper set {', '.join(f + '=?' for f in FIELDS)} "
                   f"where paper_id=?", [rec[f] for f in FIELDS] + [pid])
    pe.commit()
    print(f"database updated to match ({len(want)} rows)")


if __name__ == "__main__":
    main()
