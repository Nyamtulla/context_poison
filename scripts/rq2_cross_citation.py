"""Regenerate cross_citation_analysis.md from the current Excel corpus state.

RQ2: do Track A (Security) and Track B (ML/AI) papers cite each other, or
have they evolved as disconnected communities? Computed over the full
population (not a stratified sample) using the already-computed
cites_track_a/cites_track_b columns, which are mechanically derived from the
real citation graph -- see the "Methodology note" in the generated report
for what this does and doesn't verify.

Usage:
    python3 scripts/rq2_cross_citation.py
"""
import datetime
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

EXCEL_PATH = REPO_ROOT / "data" / "exports" / "paper_dashboard_source.xlsx"
OUTPUT_PATH = REPO_ROOT / "cross_citation_analysis.md"


def load_corpus():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb["Papers"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[idx["screening"]] != "Exclude"]
    wb.close()
    return rows, idx


def compute_rates(rows, idx):
    track_a = track_b = both = 0
    a_cites_b = b_cites_a = 0
    grade_stats = {"Security": {}, "ML/AI": {}}
    for row in rows:
        track = row[idx["track"]]
        grade = row[idx["evidence_grade"]] or "?"
        ca = row[idx["cites_track_a"]]
        cb = row[idx["cites_track_b"]]
        if track == "Security":
            track_a += 1
            grade_stats["Security"].setdefault(grade, [0, 0])
            grade_stats["Security"][grade][0] += 1
            if cb == "Y":
                a_cites_b += 1
                grade_stats["Security"][grade][1] += 1
        elif track == "ML/AI":
            track_b += 1
            grade_stats["ML/AI"].setdefault(grade, [0, 0])
            grade_stats["ML/AI"][grade][0] += 1
            if ca == "Y":
                b_cites_a += 1
                grade_stats["ML/AI"][grade][1] += 1
        elif track == "Both":
            both += 1
    return {
        "track_a": track_a, "track_b": track_b, "both": both,
        "a_cites_b": a_cites_b, "b_cites_a": b_cites_a,
        "grade_stats": grade_stats,
    }


def render_report(stats):
    ta, tb = stats["track_a"], stats["track_b"]
    ab, ba = stats["a_cites_b"], stats["b_cites_a"]
    lines = []
    lines.append("# Cross-Citation Analysis (RQ2)")
    lines.append("")
    lines.append(f"Regenerated {datetime.date.today().isoformat()} by `scripts/rq2_cross_citation.py`.")
    lines.append("")
    lines.append("**RQ2:** Do the agent-security-poisoning literature (Track A) and the "
                  "agentic-context-management/reliability literature (Track B) cite each other, "
                  "or have they evolved as disconnected communities studying the same failure mode?")
    lines.append("")
    lines.append("## Headline result")
    lines.append("")
    lines.append(f"Computed over the **entire coded population** ({ta + tb} papers classified as "
                  f"Track A/Security or Track B/ML-AI; {stats['both']} additional papers classified "
                  "as \"Both\"), not a sample:")
    lines.append("")
    lines.append("| Direction | Papers citing >=1 paper from the other track | Total papers in track | Rate |")
    lines.append("|---|---|---|---|")
    lines.append(f"| Track A (Security) -> cites Track B (ML/AI) | {ab} | {ta} | **{ab/ta*100:.1f}%** |")
    lines.append(f"| Track B (ML/AI) -> cites Track A (Security) | {ba} | {tb} | **{ba/tb*100:.1f}%** |")
    lines.append("")
    lines.append("Because this is a full-population count (every paper's citation edges are "
                  "resolved mechanically from the actual citation graph, not hand-sampled), no "
                  "confidence interval is reported.")
    lines.append("")
    lines.append("## By evidence grade")
    lines.append("")
    lines.append("| Grade | Track A citing B | Track B citing A |")
    lines.append("|---|---|---|")
    grades = sorted(set(stats["grade_stats"]["Security"]) | set(stats["grade_stats"]["ML/AI"]))
    for g in grades:
        a_tot, a_hit = stats["grade_stats"]["Security"].get(g, [0, 0])
        b_tot, b_hit = stats["grade_stats"]["ML/AI"].get(g, [0, 0])
        a_str = f"{a_hit}/{a_tot} = {a_hit/a_tot*100:.1f}%" if a_tot else "n/a"
        b_str = f"{b_hit}/{b_tot} = {b_hit/b_tot*100:.1f}%" if b_tot else "n/a"
        lines.append(f"| {g} | {a_str} | {b_str} |")
    lines.append("")
    lines.append("## Methodology note")
    lines.append("")
    lines.append("`cites_track_a`/`cites_track_b` are computed mechanically from the real "
                  "citation graph against the automated `track_auto`/`track_human` classification "
                  "of every cited paper in the full underlying corpus (not just the currently-"
                  "included working set) -- i.e., a paper is scored as \"cites Track A\" if it "
                  "cites *any* paper anywhere in the graph that resolved to Track A. This is "
                  "broader-coverage than a curated keyword+author signature list but inherits "
                  "whatever noise exists in the automated track classification. Manual "
                  "verification of a stratified sample against the automated classification is "
                  "recommended before treating this as a final number for publication; see "
                  "`context_integrity_sok_project_plan.md` Section 12 (threats to validity).")
    lines.append("")
    return "\n".join(lines)


def main():
    rows, idx = load_corpus()
    stats = compute_rates(rows, idx)
    report = render_report(stats)
    OUTPUT_PATH.write_text(report)
    print(f"Wrote {OUTPUT_PATH}")
    print(f"A->B: {stats['a_cites_b']}/{stats['track_a']} = {stats['a_cites_b']/stats['track_a']*100:.1f}%")
    print(f"B->A: {stats['b_cites_a']}/{stats['track_b']} = {stats['b_cites_a']/stats['track_b']*100:.1f}%")


if __name__ == "__main__":
    main()
