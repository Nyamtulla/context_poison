"""Split the validation master into one workbook per coder.

Run when you know how many people you have:

    python3 scripts/split_validation_packets.py --coders 4

Two things this is careful about.

**Every coder gets a mix of all three confidence groups.** It deals the rows
round-robin from the master's already-shuffled order rather than slicing it
into blocks. A coder who happened to get mostly low-confidence entries would
return a low accuracy that says more about their packet than about the
extraction, and if that coder is also stricter than the others the two effects
are impossible to separate afterwards.

**No coder can tell which group a row is in.** The confidence column is not in
the master and is not added here. The deal order is the only thing that could
leak it, and round-robin over a shuffled list leaks nothing.

Re-running with the same `--coders` gives the same split; changing the number
reshuffles everything, so decide the headcount before sending anything out.

Out: data/exports/confidence_validation/packets/validation-coder-NN.xlsx
"""
from __future__ import annotations

import argparse
import collections
import json
import pathlib
import shutil

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO = pathlib.Path(__file__).resolve().parent.parent
SRC = REPO / "data" / "exports" / "confidence_validation"
OUT = SRC / "packets"

FILLCOLS = {"quote_supports", "name_correct", "class_correct", "verdict",
            "correct_value", "minutes", "notes"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--coders", type=int, required=True)
    args = ap.parse_args()
    if args.coders < 1:
        raise SystemExit("need at least one coder")

    master = SRC / "validation_master.xlsx"
    if not master.exists():
        raise SystemExit(f"no {master.name}; run build_confidence_validation.py")
    wb = openpyxl.load_workbook(master)
    grid = list(wb["check"].iter_rows(values_only=True))
    hdr, rows = list(grid[0]), [list(r) for r in grid[1:] if r and r[0]]

    # round-robin, not blocks: see the note at the top
    deal = collections.defaultdict(list)
    for i, r in enumerate(rows):
        deal[i % args.coders].append(r)

    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    key = json.loads((SRC / "_key_DO_NOT_SHARE.json").read_text())["key"]
    intro = [c.value for c in
             next(wb["start here"].iter_cols(min_col=1, max_col=1))]

    for c in sorted(deal):
        out = openpyxl.Workbook()
        ws = out.active
        ws.title = "start here"
        for i, line in enumerate(intro):
            ws.cell(row=i + 1, column=1, value=line)
        ws.cell(row=len(intro) + 2, column=1,
                value=f"This is packet {c+1} of {args.coders}. "
                      f"{len(deal[c])} rows, roughly "
                      f"{round(len(deal[c])*8/60, 1)} hours at 8 minutes each.")
        ws.column_dimensions["A"].width = 100

        cs = out.create_sheet("check")
        cs.append(hdr)
        yellow = PatternFill("solid", fgColor="FFF2CC")
        grey = PatternFill("solid", fgColor="E8E8E8")
        for n, h in enumerate(hdr, 1):
            cell = cs.cell(row=1, column=n)
            cell.font = Font(bold=True)
            cell.fill = yellow if h in FILLCOLS else grey
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in deal[c]:
            cs.append(r)
        widths = {"paper_title": 46, "evidence_quote": 70, "recorded_name": 30,
                  "pdf": 34, "correct_value": 30, "notes": 34, "entry_id": 14}
        for n, h in enumerate(hdr, 1):
            cs.column_dimensions[cs.cell(row=1, column=n).column_letter].width \
                = widths.get(h, 13)
        for r in range(2, len(deal[c]) + 2):
            for h in ("evidence_quote", "paper_title"):
                cs.cell(row=r, column=hdr.index(h) + 1).alignment = \
                    Alignment(wrap_text=True, vertical="top")
        cs.freeze_panes = "A2"
        for col, opts in (("quote_supports", '"yes,partly,no"'),
                          ("name_correct", '"yes,no"'),
                          ("class_correct", '"yes,no"'),
                          ("verdict", '"correct,minor,wrong"')):
            dv = DataValidation(type="list", formula1=opts, allow_blank=True)
            cs.add_data_validation(dv)
            L = cs.cell(row=1, column=hdr.index(col) + 1).column_letter
            dv.add(f"{L}2:{L}{len(deal[c])+1}")

        f = OUT / f"validation-coder-{c+1:02d}.xlsx"
        out.save(f)
        mix = collections.Counter(key.get(r[0], "?") for r in deal[c])
        print(f"  {f.name}: {len(deal[c]):>3} rows   "
              f"high {mix['high']:>2}  medium {mix['medium']:>2}  "
              f"low {mix['low']:>2}   ~{len(deal[c])*8/60:.1f} h")

    print(f"\n{len(rows)} rows across {args.coders} coders -> {OUT}")
    print("Send the packets. Do NOT send _key_DO_NOT_SHARE.json.")


if __name__ == "__main__":
    main()
