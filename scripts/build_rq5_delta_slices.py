"""Prepare RQ5 coverage-matching work slices for newly-registered defenses.

After build_registry.py rq4 runs, some defenses in the registry are new (they
came from the screening-delta extraction) and have never been matched against
the RQ3 mechanism registry. This script emits:

  * data/registries/raw/rq3_registry_names.json  -- the COMPLETE current RQ3
    registry, the lookup reference every matching agent needs;
  * data/registries/raw/rq5_delta_slice<N>.json  -- the new defenses, split into
    N slices, each carrying the source paper's own evaluation text.

Only defenses with no existing rq5 pair AND no prior match attempt are included,
so re-running is cheap and does not re-litigate settled pairs.

    python3 scripts/build_rq5_delta_slices.py [--slices 4]
"""
from __future__ import annotations
import argparse, csv, json, sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
RAW = REPO / "data/registries/raw"
REG = REPO / "data/registries"
sys.path.insert(0, str(REPO))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--slices", type=int, default=4)
    args = ap.parse_args()

    mechs = json.loads((REG / "rq3_pollution_registry.json").read_text())
    defs_ = json.loads((REG / "rq4_defense_registry.json").read_text())

    # lookup reference for the agents: every mechanism name, verbatim
    ref = [{"mechanism_name": m["name"], "channel": m.get("channel"),
            "consequence": m.get("consequence"), "from_paper": m.get("title", "")[:90]}
           for m in mechs]
    (RAW / "rq3_registry_names.json").write_text(json.dumps(ref, indent=1) + "\n")
    print(f"rq3_registry_names.json: {len(ref)} mechanisms")

    # which defense rows already have a match attempt in any rq5 batch?
    attempted = set()
    for fp in list(RAW.glob("rq5_batch*.csv")) + list(RAW.glob("rq5_delta_batch*.csv")):
        with open(fp) as f:
            for r in csv.DictReader(f):
                try:
                    attempted.add(int(r["defense_row"]))
                except (KeyError, ValueError):
                    pass
    print(f"defense rows with an existing match attempt: {len(attempted)}")

    # paper text for the defenses that still need matching
    import openpyxl
    wb = openpyxl.load_workbook(REPO / "data/exports/paper_dashboard_source.xlsx", read_only=True)
    ws = wb["Papers"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    i = {h: n for n, h in enumerate(hdr)}
    text = {}
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        text[n] = {k: row[i[k]] for k in
                   ("title", "baselines_compared", "key_result",
                    "technical_summary", "datasets_benchmarks")}
    wb.close()

    todo = []
    for d in defs_:
        row = int(d["row"])
        if row in attempted:
            continue
        t = text.get(row, {})
        todo.append({"defense_row": row, "defense_name": d["name"],
                     "source_title": t.get("title", ""),
                     "baselines_compared": t.get("baselines_compared", ""),
                     "key_result": t.get("key_result", ""),
                     "technical_summary": t.get("technical_summary", ""),
                     "datasets_benchmarks": t.get("datasets_benchmarks", "")})
    print(f"defenses needing a match pass: {len(todo)}")
    if not todo:
        print("nothing to do.")
        return

    n = max(1, min(args.slices, len(todo)))
    size = -(-len(todo) // n)
    for s in range(n):
        chunk = todo[s * size:(s + 1) * size]
        if not chunk:
            continue
        fp = RAW / f"rq5_delta_slice{s+1}.json"
        fp.write_text(json.dumps(chunk, indent=1) + "\n")
        print(f"  {fp.name}: {len(chunk)} defenses")


if __name__ == "__main__":
    main()
