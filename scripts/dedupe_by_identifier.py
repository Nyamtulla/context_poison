"""Third-pass corpus dedup: same paper, same identifier, different title.

The cheapest and strongest duplicate signal in this corpus is the identifier,
and until 2026-09-21 nothing checked it. Two papers can be the same work under
different titles, but they cannot have the same DOI or the same arXiv number by
accident.

This catches what the other two passes structurally cannot:

  * `dedupe_corpus.py` matches TITLES, so a preprint renamed on publication
    slips through (AgentFuzzer/AgentVigil, InjecGuard/PIGuard).
  * `dedupe_by_abstract.py` matches ABSTRACTS, so it skips any row whose
    abstract is missing -- which is exactly how the Liu et al. Open-Prompt-
    Injection pair survived it: row 1126 has an empty abstract, so there was
    nothing to compare.

The crucial normalisation is that an arXiv id hides in two places. Some rows
carry it in `arxiv_id`; others carry only a DOI of the form
`10.48550/arXiv.2310.12815`, which is the same identifier wearing a hat. Rows
1011 and 1126 are one paper precisely because one's `arxiv_id` equals the
other's DOI suffix.

Report-only. Confirm each group, then exclude via dedupe_corpus.py --apply.

    python3 scripts/dedupe_by_identifier.py
"""
from __future__ import annotations
import re, sys
from collections import defaultdict
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
XL = REPO / "data/exports/paper_dashboard_source.xlsx"


def norm_arxiv(arxiv_id, doi):
    """Return a canonical arXiv number from either field, or None."""
    for raw in (arxiv_id, doi):
        if not raw:
            continue
        s = str(raw).strip().lower()
        s = s.replace("arxiv:", "")
        m = re.search(r"10\.48550/arxiv\.(\S+)", s)
        if m:
            s = m.group(1)
        m = re.fullmatch(r"(\d{4}\.\d{4,5})(v\d+)?", s)
        if m:
            return m.group(1)
    return None


def norm_doi(doi):
    if not doi:
        return None
    s = str(doi).strip().lower().replace("https://doi.org/", "")
    if s.startswith("10.48550/arxiv."):
        return None          # that is an arXiv id, handled above
    return s if s.startswith("10.") else None


def main() -> None:
    wb = openpyxl.load_workbook(XL, read_only=True)
    ws = wb["Papers"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    i = {h: n for n, h in enumerate(hdr)}

    by_arxiv, by_doi, rows = defaultdict(list), defaultdict(list), {}
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if r[i["screening"]] == "Exclude":
            continue
        rows[n] = r
        a = norm_arxiv(r[i["arxiv_id"]], r[i["doi"]])
        d = norm_doi(r[i["doi"]])
        if a:
            by_arxiv[a].append(n)
        if d:
            by_doi[d].append(n)
    wb.close()

    found = 0
    for label, index in (("arXiv id", by_arxiv), ("DOI", by_doi)):
        groups = {k: v for k, v in index.items() if len(v) > 1}
        if not groups:
            print(f"No duplicate {label} groups among included papers.")
            continue
        print(f"\n{len(groups)} duplicate {label} group(s):\n")
        for key, members in sorted(groups.items()):
            found += 1
            print(f"  {label} {key}")
            for n in members:
                r = rows[n]
                print(f"    row {n:>5}  {r[i['year']]}  {r[i['citation_count']] or 0:>4}c  "
                      f"{str(r[i['venue']] or '')[:22]:<22} abs={len(r[i['abstract']] or '')}")
                print(f"           {str(r[i['title']])[:86]}")
            print()
    if not found:
        print("\nCorpus is clean on the identifier axis.")
    else:
        print("Confirm each group, then: python3 scripts/dedupe_corpus.py --apply <row>,...")


if __name__ == "__main__":
    main()
