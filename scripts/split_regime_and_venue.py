"""Make regime and venue two separate columns, because they are two things.

Until 2026-10-22 one column, `track`, was doing three jobs: which literature
a paper belongs to, which threat model it answers to, and - after the venue
rule landed - where it published. Collapsing those produced a transfer matrix
that counted ARGUS against BIPIA as a cross-community pair, when BIPIA is a
prompt-injection benchmark that merely published at an AI venue.

So:

  regime  security | safety   is an adversary involved, or does the context
                              go wrong on its own. A property of the work.
                              Restored from the full-text re-coding, which
                              judged it from the paper itself.

  venue   security | AI | -   where the paper published. Observable. Derived
                              by scripts/retrack_by_venue.py; arXiv and
                              unrecognised venues carry "-".

Every taxonomy and transfer claim is about REGIME. Venue is its own finding:
it says where each kind of work goes, and whether the two kinds cite across.
"""
import collections, datetime, importlib.util, pathlib, shutil
import openpyxl

_spec = importlib.util.spec_from_file_location(
    'rtv', pathlib.Path(__file__).with_name('retrack_by_venue.py'))
_rtv = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(_rtv)

XL = "data/exports/paper_dashboard_source.xlsx"
BACKUP = "data/exports/paper_dashboard_source.backup-venue-20261022-170101.xlsx"
REGIME = {"Security": "security", "ML/AI": "safety"}


def main():
    pre = openpyxl.load_workbook(BACKUP, read_only=True)["Papers"]
    rows = list(pre.iter_rows(values_only=True))
    h = {k: i for i, k in enumerate(rows[0]) if k}
    content = {r[h["paper_id"]]: str(r[h["track"]] or "").strip()
               for r in rows[1:]}

    wb = openpyxl.load_workbook(XL)
    ws = wb["Papers"]
    hdr = [c.value for c in ws[1]]
    ti, pi, si, vni = (hdr.index(x) + 1
                       for x in ("track", "paper_id", "screening", "venue"))

    # two new columns rather than overloading one
    for name in ("regime", "venue_community"):
        if name not in hdr:
            hdr.append(name)
            ws.cell(1, len(hdr)).value = name
    ri, vi = hdr.index("regime") + 1, hdr.index("venue_community") + 1

    tally = collections.Counter()
    for row in range(2, ws.max_row + 1):
        pid = str(ws.cell(row, pi).value or "").strip()
        # derive venue from the venue STRING, never from the track column,
        # which the venue pass already overwrote for some rows
        vt, _why = _rtv.bucket(str(ws.cell(row, vni).value or ""))
        ws.cell(row, vi).value = {"Security": "security",
                                  "ML/AI": "AI"}.get(vt, "-")
        back = content.get(pid, "")
        ws.cell(row, ri).value = REGIME.get(back, "")
        # `track` goes back to the content judgement, which is what every
        # downstream script means by it.
        ws.cell(row, ti).value = back
        if str(ws.cell(row, si).value or "").strip() != "Exclude":
            tally[(REGIME.get(back, "?"), ws.cell(row, vi).value)] += 1

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy(XL, XL.replace(".xlsx", f".backup-split-{stamp}.xlsx"))
    wb.save(XL)

    print(f"{'':10} " + "".join(f"{v:>12}" for v in ("security", "AI", "-")))
    for r in ("security", "safety"):
        print(f"{r:10} " + "".join(f"{tally[(r, v)]:>12}"
                                   for v in ("security", "AI", "-")))
    print("\nrows are REGIME (is there an adversary), columns are VENUE")


if __name__ == "__main__":
    main()
