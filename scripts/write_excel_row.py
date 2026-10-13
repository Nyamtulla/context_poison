import sys
import json
import sqlite3
import pathlib
import openpyxl

REPO = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "scripts"))

PATH = "/home/n646s681/context_sok/data/exports/paper_dashboard_source.xlsx"

def title_to_filename(title):
    """Same convention as scripts/rename_pdfs_to_titles.py."""
    from rename_pdfs_to_titles import sanitize
    return sanitize(title, "unknown") + ".pdf"


def write_rows(records):
    wb = openpyxl.load_workbook(PATH)
    ws = wb["Papers"]
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) + 1 for name in header}
    for rec in records:
        row = rec.pop("row")
        for field, value in rec.items():
            ws.cell(row=row, column=idx[field], value=value)
        # mark PDF path now that it's downloaded.
        # Corpus PDFs are named by TITLE, not by arXiv id - see
        # scripts/rename_pdfs_to_titles.py. This used to write
        # data/papers/{arxiv_id}.pdf, which stopped resolving after the rename.
        # Prefer the path the database already holds; only fall back to a
        # constructed name when the paper is not in the database yet.
        pid = ws.cell(row=row, column=idx["paper_id"]).value if "paper_id" in idx else None
        path = None
        if pid:
            con = sqlite3.connect(REPO / "data" / "paper_explorer.db")
            got = con.execute(
                "select pdf_local_path from paper where paper_id=?", (pid,)).fetchone()
            con.close()
            path = got[0] if got and got[0] else None
        if path is None:
            title = ws.cell(row=row, column=idx["title"]).value if "title" in idx else None
            if title:
                path = f"data/papers/{title_to_filename(title)}"
        if path:
            ws.cell(row=row, column=idx["pdf_local_path"], value=path)
    wb.save(PATH)
    print(f"wrote {len(records)} rows")

if __name__ == "__main__":
    records = json.loads(sys.stdin.read())
    write_rows(records)
