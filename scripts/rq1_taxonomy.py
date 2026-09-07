"""Regenerate rq1_taxonomy_analysis.md from the current Excel corpus state.

RQ1: across channel x intent (track) x consequence, which cells have been
studied and which are empty? Also builds the channel x defense-intervention-
point coverage matrix (the "pond" deliverable).

Purely mechanical -- channel/consequence/track/defense_intervention_point
are already-coded columns, this just pivots them. Safe and cheap to re-run
any time the corpus (screening decisions, coding) changes.

Usage:
    python3 scripts/rq1_taxonomy.py
"""
import datetime
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

EXCEL_PATH = REPO_ROOT / "data" / "exports" / "paper_dashboard_source.xlsx"
OUTPUT_PATH = REPO_ROOT / "rq1_taxonomy_analysis.md"

CHANNELS = ["memory", "RAG", "tool-output", "tool-metadata", "skill",
            "multi-agent", "cross-modal", "supply-chain", "direct-input"]
TRACKS = ["Security", "ML/AI", "Both"]
CONSEQUENCES = ["goal-hijack", "data-exfiltration", "persistence-backdoor",
                "resource-abuse", "silent-corruption", "reasoning-corruption"]
DEFENSE_POINTS = ["ingestion", "reasoning", "execution", "none"]


def load_corpus():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb["Papers"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[idx["screening"]] != "Exclude"]
    wb.close()
    return rows, idx


def build_cube(rows, idx):
    cube = defaultdict(int)
    for row in rows:
        ch, tr, co = row[idx["channel"]], row[idx["track"]], row[idx["consequence"]]
        if ch and tr and co:
            cube[(ch, tr, co)] += 1
    return cube


def build_defense_matrix(rows, idx):
    matrix = defaultdict(int)
    for row in rows:
        ch, dp = row[idx["channel"]], row[idx["defense_intervention_point"]]
        if ch and dp:
            matrix[(ch, dp)] += 1
    return matrix


def render_report(rows, idx, cube, matrix):
    total = len(rows)
    total_cells = len(CHANNELS) * len(TRACKS) * len(CONSEQUENCES)
    empty_cells = [(ch, tr, co) for ch in CHANNELS for tr in TRACKS for co in CONSEQUENCES if cube[(ch, tr, co)] == 0]
    empty_pct = len(empty_cells) / total_cells * 100

    ch_totals = defaultdict(int)
    for row in rows:
        if row[idx["channel"]]:
            ch_totals[row[idx["channel"]]] += 1

    co_by_track = defaultdict(lambda: defaultdict(int))
    for row in rows:
        tr, co = row[idx["track"]], row[idx["consequence"]]
        if tr and co:
            co_by_track[tr][co] += 1

    lines = []
    lines.append("# RQ1 -- Taxonomy Analysis: channel x intent x consequence")
    lines.append("")
    lines.append(f"Regenerated {datetime.date.today().isoformat()} by `scripts/rq1_taxonomy.py` "
                  f"from the current Excel corpus state ({total} included papers). "
                  "Re-run this script any time screening/coding changes.")
    lines.append("")
    lines.append("**RQ1:** Across the full landscape of LLM agent context contamination -- "
                  "memory, RAG, tool output, tool metadata, skills, multi-agent, cross-modal, "
                  "supply chain -- which channel x intent (adversarial/incidental) x consequence "
                  "cells have been studied, and which are empty?")
    lines.append("")
    lines.append("## Headline result")
    lines.append("")
    lines.append(f"**{empty_pct:.1f}% of the channel x intent x consequence cube is empty** "
                  f"({len(empty_cells)} of {total_cells} cells: {len(CHANNELS)} channels x "
                  f"{len(TRACKS)} intents x {len(CONSEQUENCES)} consequences).")
    lines.append("")
    lines.append("## Channel totals (both tracks combined)")
    lines.append("")
    lines.append("| Channel | Papers | Share |")
    lines.append("|---|---:|---:|")
    for ch in sorted(CHANNELS, key=lambda c: -ch_totals[c]):
        lines.append(f"| {ch} | {ch_totals[ch]} | {ch_totals[ch]/total*100:.1f}% |")
    lines.append("")
    lines.append("## Consequence totals by track (intent)")
    lines.append("")
    lines.append("| Consequence | Security (Track A) | ML/AI (Track B) | Both |")
    lines.append("|---|---:|---:|---:|")
    for co in CONSEQUENCES:
        lines.append(f"| {co} | {co_by_track['Security'][co]} | {co_by_track['ML/AI'][co]} | {co_by_track['Both'][co]} |")
    lines.append("")
    lines.append("## Channel x Defense-intervention-point coverage")
    lines.append("")
    lines.append("| Channel | " + " | ".join(dp.capitalize() for dp in DEFENSE_POINTS) + " |")
    lines.append("|---|" + "---:|" * len(DEFENSE_POINTS))
    for ch in CHANNELS:
        vals = [str(matrix[(ch, dp)]) for dp in DEFENSE_POINTS]
        lines.append(f"| {ch} | " + " | ".join(vals) + " |")
    lines.append("")
    lines.append("## Full 3D cube")
    lines.append("")
    for tr in TRACKS:
        lines.append(f"### {tr}")
        lines.append("")
        lines.append("| channel | " + " | ".join(co for co in CONSEQUENCES) + " |")
        lines.append("|---|" + "---:|" * len(CONSEQUENCES))
        for ch in CHANNELS:
            vals = [str(cube[(ch, tr, co)]) for co in CONSEQUENCES]
            lines.append(f"| {ch} | " + " | ".join(vals) + " |")
        lines.append("")
    lines.append("## Empty cells by channel")
    lines.append("")
    by_channel = defaultdict(list)
    for ch, tr, co in empty_cells:
        by_channel[ch].append((tr, co))
    for ch in CHANNELS:
        lines.append(f"- **{ch}**: {len(by_channel[ch])} of {len(TRACKS)*len(CONSEQUENCES)} empty")
    lines.append("")
    lines.append("## Methodology / caveats")
    lines.append("")
    lines.append("- `channel`, `consequence`, `track`, and `defense_intervention_point` are "
                  "already-coded Excel columns (Week-2 batch classification); this script is a "
                  "pure pivot over them, no new judgment is applied here.")
    lines.append("- A corpus re-screening/deduplication pass should be run before this script "
                  "if the underlying `screening` column may be stale -- see "
                  "`scripts/dedupe_corpus.py` and `scripts/rescreen_corpus.py`.")
    lines.append("")

    return "\n".join(lines)


def main():
    rows, idx = load_corpus()
    cube = build_cube(rows, idx)
    matrix = build_defense_matrix(rows, idx)
    report = render_report(rows, idx, cube, matrix)
    OUTPUT_PATH.write_text(report)
    print(f"Wrote {OUTPUT_PATH} ({len(rows)} papers)")


if __name__ == "__main__":
    main()
