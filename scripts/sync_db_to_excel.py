"""Push corpus-state changes from the SQLite DB back into the authoritative Excel.

`data/exports/paper_dashboard_source.xlsx` is the source of truth:
`build_paper_db.py` rebuilds `paper_explorer.db` FROM it, and the 15 full-text
analysis columns exist nowhere else. Edits made directly in the DB - screening
decisions, pdf_local_path - are therefore invisible to a rebuild and are
silently reverted by one.

Today's work was done in the DB: 42 screening changes and 1,127 pdf paths
(every PDF was renamed to its title, and 30 new ones were registered). Without
this sync, `rebuild-corpus` would restore the corpus to 1,026 and throw all of
it away.

Writes a timestamped backup of the workbook before touching it.
"""
from __future__ import annotations

import pathlib
import shutil
import sqlite3
import sys
import time

import openpyxl

REPO = pathlib.Path(__file__).resolve().parent.parent
XL = REPO / "data" / "exports" / "paper_dashboard_source.xlsx"
DB = REPO / "data" / "paper_explorer.db"
SYNC = ["screening", "pdf_local_path"]


def main() -> None:
    dry = "--apply" not in sys.argv
    pe = sqlite3.connect(DB)
    db = {p: dict(zip(SYNC, rest))
          for p, *rest in pe.execute(f"select paper_id, {', '.join(SYNC)} from paper")}

    wb = openpyxl.load_workbook(XL)
    ws = wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[1]]
    col = {h: n + 1 for n, h in enumerate(hdr)}
    missing = [c for c in SYNC + ["paper_id"] if c not in col]
    if missing:
        raise SystemExit(f"workbook lacks columns: {missing}")

    changes = []
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(row=r, column=col["paper_id"]).value
        if not pid or pid not in db:
            continue
        for f in SYNC:
            old = ws.cell(row=r, column=col[f]).value
            new = db[pid][f]
            if (str(old or "").strip()) != (str(new or "").strip()):
                changes.append((r, f, old, new))
                if not dry:
                    ws.cell(row=r, column=col[f], value=new)

    by = {}
    for _, f, _, _ in changes:
        by[f] = by.get(f, 0) + 1
    print(f"{len(changes)} cell changes: {by}")
    if dry:
        print("\ndry run; pass --apply to write")
        return
    bak = XL.with_suffix(f".backup-{time.strftime('%Y%m%d-%H%M%S')}.xlsx")
    shutil.copy2(XL, bak)
    wb.save(XL)
    print(f"backed up -> {bak.name}")
    print(f"wrote {XL.relative_to(REPO)}")


if __name__ == "__main__":
    main()
